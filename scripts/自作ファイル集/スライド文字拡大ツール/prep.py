# -*- coding: utf-8 -*-
"""対象PDFの頁→pptxのスライドを突き合わせ、relayout をかけ、前後比較の画像を作る"""
import os, re, sys, glob, difflib, shutil, json, fitz
from pptx import Presentation
from win32com.client import gencache
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import relayout
H = os.path.dirname(os.path.abspath(__file__))
BASE = r"C:\Users\junch\OneDrive\_旧デスクトップ"
def find_pptx(pdf):
    stem = os.path.splitext(os.path.basename(pdf))[0]
    c = [f for f in glob.glob(BASE + r"\**\%s.pptx" % glob.escape(stem), recursive=True) if "\_" not in f.replace(BASE, "").rsplit("\\", 1)[0].replace("\_スライド版", "").replace("\_編集用原本（Word・PPT）", "")]
    return c
def stext(sl):
    out = []
    def rec(s):
        if s.shape_type == 6:
            for c in s.shapes: rec(c)
        elif s.has_text_frame: out.append(s.text_frame.text)
        elif getattr(s, "has_table", False) and s.has_table:
            for r in s.table.rows:
                for c in r.cells: out.append(c.text)
    for s in sl.shapes: rec(s)
    return re.sub(r"\s", "", "".join(out))
def mapping(pdf, pptx, pages):
    d = fitz.open(pdf); pr = Presentation(pptx); st = [stext(s) for s in pr.slides]; m = {}
    for pn in pages:
        t = re.sub(r"\s", "", d[pn - 1].get_text())
        best = max(range(len(st)), key=lambda i: difflib.SequenceMatcher(None, st[i][:400], t[:400]).ratio())
        m[pn] = (best + 1, round(difflib.SequenceMatcher(None, st[best][:400], t[:400]).ratio(), 2))
    return m
def run(pdf, pages, tag, fmax=2.0):
    cands = find_pptx(pdf); pptx = cands[0]
    m = mapping(pdf, pptx, pages)
    src = os.path.join(H, tag + "_src.pptx"); dst = os.path.join(H, tag + "_t.pptx"); shutil.copyfile(pptx, src)
    sl = sorted(set(v[0] for v in m.values()))
    res = relayout.main(src, dst, sl, fmax)
    app = gencache.EnsureDispatch("PowerPoint.Application"); from PIL import Image
    ims = []
    for f in (src, dst):
        pr = app.Presentations.Open(f, True, False, False)
        for n in sl:
            p = os.path.join(H, "%s_%s_%d.png" % (tag, "a" if f == src else "b", n)); pr.Slides(n).Export(p, "PNG", 640, 360); ims.append(Image.open(p))
        pr.Close()
    k = len(sl); s = Image.new("RGB", (2 * 644, k * 364), "gray")
    for i in range(k):
        s.paste(ims[i], (0, i * 364)); s.paste(ims[k + i], (644, i * 364))
    s.save(os.path.join(H, tag + "_cmp.png"))
    json.dump({"pdf": pdf, "pptx": pptx, "map": {str(a): b for a, b in m.items()}, "res": res}, open(os.path.join(H, tag + ".json"), "w", encoding="utf-8"), ensure_ascii=False)
    return pptx, m, res
if __name__ == "__main__":
    print(run(sys.argv[1], [int(x) for x in sys.argv[2].split(",")], sys.argv[3]))
