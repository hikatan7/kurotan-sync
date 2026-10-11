# -*- coding: utf-8 -*-
"""手術室オカレンス 会議配布用サマリー（2ページ） 2026年8月度"""
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import rcParams
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

rcParams["font.family"] = "Yu Gothic"
rcParams["axes.unicode_minus"] = False

SCR = os.path.dirname(os.path.abspath(__file__))
BASE = r"C:\Users\junch\OneDrive\デスクトップ\自作ファイル集\オカレンスレポート作成"
SRC = os.path.join(BASE, "手術室オカレンス・マンスリーレポート_2026年6月度配布用.docx")
OUT = os.path.join(BASE, "手術室オカレンス_会議用サマリー_2026年8月度.docx")

months = ["2025/4","2025/5","2025/6","2025/7","2025/8","2025/9","2025/10","2025/11","2025/12",
          "2026/1","2026/2","2026/3","2026/4","2026/5","2026/6","2026/7","2026/8"]
anes = [62,85,101,104,72,61,71,62,51,80,68,71,78,52,96,69,84]
ce   = [46,40,39,37,36,40,45,36,38,41,31,33,30,24,23,38,36]
ns   = [None,None,44,43,39,33,19,31,26,21,23,19,21,8,11,21,23]
tot  = [(a or 0)+(c or 0)+(n or 0) for a,c,n in zip(anes,ce,ns)]

# ---- グラフ1：職種別推移＋合計 ----
trend_png = os.path.join(SCR, "sum_aug_trend.png")
fig, ax = plt.subplots(figsize=(9.6, 2.7), dpi=220)
x = list(range(len(months)))
ax.bar(x, tot, color="#e8eef7", label="合計", width=0.62)
ax.plot(x, anes, marker="o", ms=4, lw=1.8, color="#1f4e79", label="麻酔科")
ax.plot(x, ce, marker="s", ms=4, lw=1.8, color="#c00000", label="CE")
ax.plot(x, [v if v is not None else float("nan") for v in ns], marker="^", ms=4, lw=1.8,
        color="#2e7d32", label="NS")
i = len(months)-1
ax.annotate(f"合計{tot[i]}", (i, tot[i]), textcoords="offset points", xytext=(0,4),
            ha="center", fontsize=8.5, color="#404040")
ax.set_xticks(x); ax.set_xticklabels(months, rotation=45, fontsize=8)
ax.set_ylabel("報告件数", fontsize=9); ax.tick_params(axis="y", labelsize=8)
ax.set_ylim(0, max(tot)+35)
ax.grid(axis="y", ls=":", color="#cccccc")
ax.legend(fontsize=8.5, ncol=4, loc="upper center", frameon=False)
for s in ("top","right"): ax.spines[s].set_visible(False)
fig.tight_layout(); fig.savefig(trend_png); plt.close(fig)

# ---- グラフ2：手術件数と100手術あたり報告件数 ----
rate_png = os.path.join(SCR, "sum_aug_rate.png")
rm = ["2026/4","2026/5","2026/6","2026/7","2026/8"]
ops = [531,448,545,617,524]
occ = [129,84,130,128,143]
rate = [o/s*100 for o,s in zip(occ,ops)]
fig, ax = plt.subplots(figsize=(9.6, 2.2), dpi=220)
ax.bar(rm, ops, color="#dce6f1", label="手術件数")
ax.set_ylabel("手術件数", fontsize=9); ax.tick_params(labelsize=8)
ax.set_ylim(0, max(ops)+170)
for i, v in enumerate(ops):
    ax.text(i, v-30, str(v), ha="center", va="top", fontsize=8.5, color="#404040")
ax2 = ax.twinx()
ax2.plot(rm, rate, marker="o", color="#c00000", lw=2, label="100手術あたり報告件数")
ax2.set_ylabel("100手術あたり報告件数", fontsize=9); ax2.tick_params(labelsize=8)
ax2.set_ylim(0, max(rate)+12)
for i, v in enumerate(rate):
    ax2.text(i, v+1.8, f"{v:.1f}", ha="center", fontsize=9, color="#c00000", fontweight="bold")
for s in ("top",): ax.spines[s].set_visible(False); ax2.spines[s].set_visible(False)
h1_, l1_ = ax.get_legend_handles_labels(); h2_, l2_ = ax2.get_legend_handles_labels()
ax.legend(h1_+h2_, l1_+l2_, fontsize=8.5, ncol=2, loc="upper left", frameon=False)
fig.tight_layout(); fig.savefig(rate_png); plt.close(fig)

# ---- グラフ3：分類別×職種（積み上げ横棒） ----
cat_png = os.path.join(SCR, "sum_aug_cat.png")
cats = ["機器","確認不足・準備不足","その他","体制・運用","薬剤","コミュニケーション"]
v_an = [18, 35, 11, 9, 10, 2]
v_ce = [36, 0, 0, 0, 0, 0]
v_ns = [1, 6, 4, 5, 2, 5]
fig, ax = plt.subplots(figsize=(9.6, 1.45), dpi=220)
y = list(range(len(cats)))[::-1]
ax.barh(y, v_an, color="#1f4e79", label="麻酔科", height=0.62)
ax.barh(y, v_ce, left=v_an, color="#c00000", label="CE", height=0.62)
ax.barh(y, v_ns, left=[a+c for a, c in zip(v_an, v_ce)], color="#2e7d32", label="NS", height=0.62)
for i, yy in enumerate(y):
    t = v_an[i]+v_ce[i]+v_ns[i]
    ax.text(t+1.2, yy, f"{cats[i]}　{t}件（{t/144*100:.1f}％）", va="center", fontsize=8.5)
ax.set_yticks([]); ax.set_xlim(0, 120)
ax.set_xticks([0,10,20,30,40,50])
ax.set_xlabel("件数（延べ144件）", fontsize=8.5); ax.tick_params(axis="x", labelsize=8)
ax.legend(fontsize=8.5, ncol=3, loc="lower right", frameon=False)
for sp_ in ("top","right","left"): ax.spines[sp_].set_visible(False)
fig.tight_layout(); fig.savefig(cat_png); plt.close(fig)

# ---------------- Word ----------------
doc = docx.Document(SRC)
body = doc.element.body
for child in list(body):
    if not child.tag.endswith("}sectPr"):
        body.remove(child)
sec = doc.sections[0]
sec.top_margin = Cm(1.2); sec.bottom_margin = Cm(1.2)
sec.left_margin = Cm(1.3); sec.right_margin = Cm(1.3)

def setfont(run, size=10, bold=False, color=None, name="Yu Gothic"):
    run.font.size = Pt(size); run.bold = bold; run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color: run.font.color.rgb = RGBColor.from_string(color)

def para(text="", size=10, bold=False, color=None, space_after=3, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if align: p.alignment = align
    if text: setfont(p.add_run(text), size, bold, color)
    return p

def h1(text, size=11.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    setfont(p.add_run(text), size, True, "1F4E79")
    pPr = p._element.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "8")
    bt.set(qn("w:space"), "1"); bt.set(qn("w:color"), "1F4E79")
    pbdr.append(bt); pPr.append(pbdr)
    return p

def bullet(text, size=9.5, bold=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.5); p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.line_spacing = 1.0
    setfont(p.add_run("・" + text), size, bold, color)
    return p

def borders(tbl, color="808080"):
    b = OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"4")
        e.set(qn("w:space"),"0"); e.set(qn("w:color"),color)
        b.append(e)
    tbl._tbl.tblPr.append(b)

def shade(cell, color):
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"),"clear")
    sh.set(qn("w:fill"), color); cell._tc.get_or_add_tcPr().append(sh)

def table(rows, header=True, size=9, widths=None, headfill="DDEBF7", after=2):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = doc.styles["Normal Table"]
    borders(t)
    for r in t.rows:
        r._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0.5); p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.line_spacing = 1.0
            setfont(p.add_run(str(val)), size, bold=(header and i == 0))
            if header and i == 0: shade(c, headfill)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(after)
    return t

def pict(path, cm=18.4):
    doc.add_picture(path, width=Cm(cm))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)

# ===== 1ページ目 =====
para("手術室オカレンス・サマリー（2026年8月度）", 15, True, "1F4E79", space_after=1)
para("集計対象：麻酔科／CE／NS オカレンスレポート 2026年8月分　　業務委員会 配布資料　　作成日：2026年9月7日　　※詳細版レポート（全7頁）別添",
     8.5, color="595959", space_after=5)

kpi = table([
    ["総手術件数","524 件\n（麻酔科管理 332 件）","総オカレンス件数","143 件\n（前月 128 件）",
     "100手術あたり報告件数","27.3 件\n（前月 20.7 件）"],
], header=False, size=10, widths=[2.4,3.4,3.0,2.8,3.4,3.2], after=2)
for j in (0, 2, 4):
    shade(kpi.cell(0, j), "DDEBF7")
    for r in kpi.cell(0, j).paragraphs:
        for run in r.runs: run.bold = True

table([
    ["職種","麻酔科","CE","NS","合計"],
    ["8月","84（前月69）","36（前月38）","23（前月21）","143（前月128）"],
    ["報告率","25.3％（÷麻酔科管理332）","6.9％（÷524）","4.4％（÷524）","27.3％（÷524）"],
], size=9, widths=[1.8,5.0,3.8,3.8,4.0], after=4)

h1("1. 月別推移（2025年4月〜2026年8月）")
pict(trend_png)

h1("2. 手術件数と報告件数（100手術あたり）")
pict(rate_png)

h1("3. 今月のポイント")
bullet("手術件数が617→524件（−15.1％）と減るなかで報告は128→143件（＋15件）と増加し、100手術あたり27.3件と2026年度で最も多く報告が集まった。オカレンスはアクシデントと異なり、件数の増加は現場のヒアリング文化・報告文化の定着を示すものとして肯定的に評価する。", 9.5, bold=True)
bullet("増加の中心は麻酔科（69→84件、率17.8→25.3件/100件）。「手術準備不足」27→35件は欠品が確実に拾い上げられている裏返し。ただし品目はBISモニター6件・ボールペン5件・ハサミ5件・挿管チューブ4件と7月と同じで、報告が集まっている今が定数化の効果を出す好機。", 9.5)
bullet("CEは36件と横ばいだが、うち麻酔器関連が12件（33.3％）。特にAtlanのフロー計測不良・リークテスト不合格が8/5〜8/27に7回反復し、複数事例で「回路内の水滴」が共通要因として記載されている。単発故障ではなく系統的課題。", 9.5)
bullet("NSは23件と2か月連続で2025年度水準を維持し、報告文化の再定着が続いている。個別の事象としては、左右確認後に反対側の目を消毒、5本抜歯の予定が4本で終了が患者本人の指摘で発覚。麻薬管理不備2件・針刺し3件とあわせ、件数とは切り離して重点対応する。", 9.5, bold=True)
bullet("上位2分類（機器55＋確認不足・準備不足41）で全体の66.7％。対策資源はこの2領域に集中する。", 9.5)

# ===== 2ページ目 =====
doc.add_page_break()

h1("4. 分類別 × 職種別（8月・延べ144件）")
pict(cat_png)

h1("5. Good Job / Good Catch（患者に到達する前に止められた事例）")
table([
    ["発生日","部署","内容","評価"],
    ["8/3","CE・麻酔科","Apollo麻酔器の始業点検でリークテスト不合格を検出。回路交換でも通過せず、麻酔中であったが2BのFabiusと交換して手術継続、本体はメーカーメンテへ","Good Catch"],
    ["8/13","CE・麻酔科","麻酔器のリークテスト不通過から呼気センサ部品の外れを発見し修正、テスト通過（朝の点検通過後に外れた可能性まで記録）","Good Catch"],
    ["8/14","CE・麻酔科","症例入れ替え時に麻酔器回路・フィルタの水滴を発見。低流量麻酔の影響を考慮し回路・フィルタ・アブソーバを全交換","Good Catch"],
    ["8/25","NS・形成外科","PM植込み患者の局所麻酔手術でCEへの連絡がなかったが、気付いた外回り看護師がCEに相談。ICD作動の可能性がある機種と判明し設定変更して実施","Good Catch"],
    ["8/19","CE・麻酔科","Atlanのフローバランス異常に対し、旧HCU機在庫のFabiusに切り替えて手術を継続","Good Job"],
], size=8.2, widths=[1.4,2.6,11.2,1.8])
para("→ 麻酔器の異常は今月も始業点検・症例間点検で捕捉できており、予備機への切り替えが手術継続を支えている。",
     8.8, color="1F4E79", space_after=3)

h1("6. 注目事象（安全上の重点・上位5件）")
table([
    ["発生日","部署","事象","必要な対応"],
    ["8/7","眼科","サインインで左右を確認したが、手術する側と反対の目を消毒。患者に言われて気付いた","消毒直前に術者・看護師で左右を声に出して読み合わせる手順を追加"],
    ["8/19","口腔外科","5本抜歯するところ4本抜歯で手術終了。帰室時に患者に聞かれ発覚","サインアウトに「予定と実施の突合（本数・個数・部位）」を追加"],
    ["8/3・8/4","看護部","使用後の麻薬を金庫へ返却せず一晩施錠されない場所に保管／麻薬が1人分不足と気付きながら未確認のまま帰宅","法令順守事項。時間外の返却手順と、数量不一致時の当日中の報告ルールを明文化・再教育"],
    ["8/31","心臓血管外科・カテ室","研修医が2週続けて縫合針で針刺し事故／カテ室で片付け時に針刺し事故","研修医への縫合針取扱い教育、ハンズフリーパスと鋭利物廃棄手順の徹底"],
    ["8/5〜8/27","CE・麻酔科","ドレーゲルAtlanのフロー計測不良・リークテスト不合格が7回反復。回路内・ソーダライム・呼気センサ周囲の水滴が共通要因","低流量麻酔時の結露対策。症例間の水滴確認を点検項目化し、点検周期をメーカーと協議"],
], size=8.2, widths=[1.5,2.6,7.3,5.6])
_e = doc.paragraphs[-1]              # 表直後の空段落を削除（2頁に収めるため）
_e._element.getparent().remove(_e._element)

h1("7. 今月の重点対策（9月）")
table([
    ["#","重点課題","具体策","主担当","期限"],
    ["1","手術部位（左右）の確認","サインインとは別に、消毒直前の左右読み合わせを追加。眼科は点眼・消毒トレイの左右分離も検討","眼科・手術室NS","9月中"],
    ["2","予定術式の完遂確認","サインアウトに本数・個数・部位の突合を追加し、読み上げ担当を固定","口腔外科・手術室NS","9月中"],
    ["3","麻薬管理（法令順守）","時間外返却手順の明文化と相互確認の必須化。数量不一致は当日中に責任者へ報告","手術室NS・薬剤部","9月中（最優先）"],
    ["4","針刺し防止","研修医への縫合針取扱い教育、ハンズフリーパスの運用、片付け時の鋭利物処理手順の再周知","医療安全・心外・NS","9月中"],
    ["5","麻酔器（Atlan）の水滴対策","症例間の水滴確認を点検項目へ追加。アブソーバ・フィルタの交換基準を明文化し、予備機の配置を維持","CE・麻酔科","9月中"],
    ["6","麻酔準備の欠品（35件）","反復品目（BISモニター・ハサミ・挿管チューブ等）を部屋別に定数化し、朝の残数確認を役割固定","手術室NS・麻酔科","継続（9月末）"],
], size=8.2, widths=[0.7,3.2,8.0,3.2,1.8])
_last = doc.paragraphs[-1]           # 表直後の空段落を削除（2頁に収めるため）
_last._element.getparent().remove(_last._element)

para("※麻酔科＝総計320件－オカレンスなし236件（分類別は延べ144件）。オカレンスは報告件数の増加を報告文化の定着として評価し、安全性は個々の事象の重大性で判断する。詳細は別添の詳細版（全7頁）参照。",
     7.8, color="595959", space_after=0)

doc.save(OUT)
print("saved:", OUT, os.path.getsize(OUT))
