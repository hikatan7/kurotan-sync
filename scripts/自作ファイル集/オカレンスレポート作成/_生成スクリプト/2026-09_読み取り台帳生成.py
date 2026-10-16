# -*- coding: utf-8 -*-
"""オカレンス読み取り台帳（2026年度上半期）"""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from occ_2026_data import *

OUT = sys.argv[1]
F = "Meiryo UI"
thin = Side(style="thin", color="A0A0A0")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="DDEBF7")
WARN = PatternFill("solid", fgColor="FFF2CC")

def sheet(wb, title, header, rows, widths, wrap_cols=()):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows: ws.append(list(r))
    for c in ws[1]:
        c.font = Font(name=F, bold=True, size=10); c.fill = HEAD; c.border = B
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=F, size=10); c.border = B
            c.alignment = Alignment(vertical="top", wrap_text=(c.column in wrap_cols))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    return ws

wb = Workbook(); wb.remove(wb.active)

# 月別集計
rows = []
for m in range(4, 10):
    a = ANES_TOTAL[m] - ANES_NONE[m]
    n = NS[MONTHS.index(f"2026/{m}")]
    c = CE[MONTHS.index(f"2026/{m}")]
    rows.append([f"2026年{m}月", ANES_TOTAL[m], ANES_NONE[m], a, c, n, a + c + n,
                 OPS[m] if OPS[m] else "（入力待ち）", ANES_OPS[m] if ANES_OPS[m] else "（入力待ち）"])
ws = sheet(wb, "月別集計", ["月","麻酔科 総数","うちオカレンスなし","麻酔科 実件数","CE","NS","合計","総手術件数","麻酔科管理件数"],
           rows, [12,12,16,13,8,16,8,13,14])
ws.append([]); ws.append(["※NS 5月は8件で確定（写真に写っていた6件のみ「NS_4-9月」に記載）。9/10のスキンテアはリムーバー使用ありで確定。"])
ws.cell(ws.max_row, 1).font = Font(name=F, size=10, color="C00000")
for r in ws.iter_rows(min_row=2, max_row=7):
    for c in r:
        if c.value == "（入力待ち）": c.fill = WARN

# 麻酔科 分類別
rows = []
for k in sorted(ANES_CAT):
    v = [ANES_BY_CAT[m].get(k, 0) for m in range(4, 10)]
    if sum(v) or ANES_BY_CAT_2025_H1.get(k):
        rows.append([k, ANES_CAT[k]] + v + [sum(v), ANES_BY_CAT_2025_H1.get(k, 0)])
rows.append([99, "オカレンスなし"] + [ANES_NONE[m] for m in range(4, 10)] + [sum(ANES_NONE.values()), ""])
rows.append(["", "総数"] + [ANES_TOTAL[m] for m in range(4, 10)] + [sum(ANES_TOTAL.values()), ANES_TOTAL_2025_H1])
sheet(wb, "麻酔科_分類別", ["番号","分類","4月","5月","6月","7月","8月","9月","上半期計","2025年度上半期"],
      rows, [6,40,6,6,6,6,6,6,10,14], wrap_cols=(2,))

# 麻酔科 9月の内訳
rows = [[k, ANES_CAT[k], n, t] for k, n, t in ANES_SEP_DETAIL]
rows.append(["", "その他、重要なもの", "", ANES_SEP_IMPORTANT])
sheet(wb, "麻酔科_9月内訳", ["番号","分類","件数","内容"], rows, [6,30,6,100], wrap_cols=(2,4))

# NS 全件
rows = [[f"{r[0]}月", r[1], r[2], r[3], r[4], r[5], r[6]] for r in NS_ROWS]
ws = sheet(wb, "NS_4-9月", ["月","日","部屋","診療科","内容","解決・対応","分類（くろたん判定）"], rows,
           [6,7,11,18,90,12,18], wrap_cols=(4,5))
for row in ws.iter_rows(min_row=2):
    if "原本で" in str(row[4].value):
        for c in row: c.fill = WARN

# CE 9月
rows = [list(r) for r in CE_SEP]
sheet(wb, "CE_9月", ["日付","主科","機器","時刻","部屋","内容","性質（くろたん判定）"], rows,
      [7,10,16,7,9,90,18], wrap_cols=(6,))

wb.save(OUT)
print("saved", OUT)
