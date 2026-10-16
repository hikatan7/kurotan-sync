# -*- coding: utf-8 -*-
"""手術室オカレンス 会議配布用サマリー（2ページ） 2026年9月度
使い方: python 2026-09_会議用サマリー生成.py <出力docx> <図の作業フォルダ>
手術件数は occ_2026_data.py の OPS[9]／ANES_OPS[9] に入れて再実行する。"""
import os, sys
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import rcParams
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SCR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, SCR)
from occ_2026_data import *

rcParams["font.family"] = "Yu Gothic"
rcParams["axes.unicode_minus"] = False

BASE = os.path.dirname(SCR)
SRC = os.path.join(BASE, "手術室オカレンス・マンスリーレポート_2026年6月度配布用.docx")
OUT = sys.argv[1]
WORK = sys.argv[2]

anes, ce, ns = ANES, CE, NS
ns_s = ns_count(9)
assert anes[-1] == ANES_TOTAL[9] - ANES_NONE[9] and ce[-1] == len(CE_SEP) and ns[-1] == ns_s
tot = [(a or 0)+(c or 0)+(n or 0) for a, c, n in zip(anes, ce, ns)]
T9, T8 = tot[-1], tot[-2]
ops9, aops9 = OPS[9], ANES_OPS[9]
h1avg = sum(tot[-6:]) / 6
WAIT = "入力待ち"

def r100(n, d): return f"{n/d*100:.1f}" if d else "―"

# ---- グラフ1：職種別推移＋合計 ----
trend_png = os.path.join(WORK, "sum_sep_trend.png")
fig, ax = plt.subplots(figsize=(9.6, 2.35), dpi=220)
x = list(range(len(MONTHS)))
ax.bar(x, tot, color="#e8eef7", label="合計", width=0.62)
ax.plot(x, anes, marker="o", ms=4, lw=1.8, color="#1f4e79", label="麻酔科")
ax.plot(x, ce, marker="s", ms=4, lw=1.8, color="#c00000", label="CE")
ax.plot(x, [v if v is not None else float("nan") for v in ns], marker="^", ms=4, lw=1.8,
        color="#2e7d32", label="NS")
i = len(MONTHS)-1
ax.annotate(f"合計{tot[i]}", (i, tot[i]), textcoords="offset points", xytext=(0,4),
            ha="center", fontsize=8.5, color="#404040")
ax.set_xticks(x); ax.set_xticklabels(MONTHS, rotation=45, fontsize=8)
ax.set_ylabel("報告件数", fontsize=9); ax.tick_params(axis="y", labelsize=8)
ax.set_ylim(0, max(tot)+35)
ax.grid(axis="y", ls=":", color="#cccccc")
ax.legend(fontsize=8.5, ncol=4, loc="upper center", frameon=False)
for s in ("top","right"): ax.spines[s].set_visible(False)
fig.tight_layout(); fig.savefig(trend_png); plt.close(fig)

# ---- グラフ2：手術件数と100手術あたり報告件数 ----
rate_png = os.path.join(WORK, "sum_sep_rate.png")
rm = [f"2026/{m}" for m in range(4, 10)]
ops = [OPS[m] for m in range(4, 10)]
occ = tot[-6:]
fig, ax = plt.subplots(figsize=(9.6, 1.85), dpi=220)
known = [v for v in ops if v]
ax.bar(rm, [v or 0 for v in ops], color="#dce6f1", label="手術件数")
ax.set_ylabel("手術件数", fontsize=9); ax.tick_params(labelsize=8)
ax.set_ylim(0, max(known)+170)
for k, v in enumerate(ops):
    if v: ax.text(k, v-30, str(v), ha="center", va="top", fontsize=8.5, color="#404040")
    else: ax.text(k, 40, "手術件数\n" + WAIT, ha="center", va="bottom", fontsize=8.5, color="#c00000")
ax2 = ax.twinx()
rate = [o/s*100 if s else float("nan") for o, s in zip(occ, ops)]
ax2.plot(rm, rate, marker="o", color="#c00000", lw=2, label="100手術あたり報告件数")
ax2.set_ylabel("100手術あたり報告件数", fontsize=9); ax2.tick_params(labelsize=8)
ax2.set_ylim(0, max(r for r in rate if r == r)+12)
for k, v in enumerate(rate):
    if v == v:
        ax2.text(k, v+1.8, f"{v:.1f}", ha="center", fontsize=9, color="#c00000", fontweight="bold")
for s in ("top",): ax.spines[s].set_visible(False); ax2.spines[s].set_visible(False)
h1_, l1_ = ax.get_legend_handles_labels(); h2_, l2_ = ax2.get_legend_handles_labels()
ax.legend(h1_+h2_, l1_+l2_, fontsize=8.5, ncol=2, loc="upper left", frameon=False)
fig.tight_layout(); fig.savefig(rate_png); plt.close(fig)

# ---- グラフ3：分類別×職種（積み上げ横棒） ----
cat_png = os.path.join(WORK, "sum_sep_cat.png")
nsg = ns_group(9)
cev = {g: 0 for g in GROUPS}; cev["機器"] = len(CE_SEP)
rows3 = []
for g in GROUPS:
    rows3.append((g, ANES_SEP_GROUP[g], cev[g], nsg[g]))
rows3.sort(key=lambda r: -(r[1]+r[2]+r[3]))
NOBE = sum(r[1]+r[2]+r[3] for r in rows3)
cats = [r[0] for r in rows3]
v_an = [r[1] for r in rows3]; v_ce = [r[2] for r in rows3]; v_ns = [r[3] for r in rows3]
fig, ax = plt.subplots(figsize=(9.6, 1.45), dpi=220)
y = list(range(len(cats)))[::-1]
ax.barh(y, v_an, color="#1f4e79", label="麻酔科", height=0.62)
ax.barh(y, v_ce, left=v_an, color="#c00000", label="CE", height=0.62)
ax.barh(y, v_ns, left=[a+c for a, c in zip(v_an, v_ce)], color="#2e7d32", label="NS", height=0.62)
for k, yy in enumerate(y):
    t = v_an[k]+v_ce[k]+v_ns[k]
    ax.text(t+1.2, yy, f"{cats[k]}　{t}件（{t/NOBE*100:.1f}％）", va="center", fontsize=8.5)
ax.set_yticks([]); ax.set_xlim(0, 110)
ax.set_xticks([0,10,20,30,40,50])
ax.set_xlabel(f"件数（延べ{NOBE}件）", fontsize=8.5); ax.tick_params(axis="x", labelsize=8)
ax.legend(fontsize=8.5, ncol=3, loc="lower right", frameon=False)
for sp_ in ("top","right","left"): ax.spines[sp_].set_visible(False)
fig.tight_layout(); fig.savefig(cat_png); plt.close(fig)
top2 = rows3[0][1]+rows3[0][2]+rows3[0][3] + rows3[1][1]+rows3[1][2]+rows3[1][3]

# ---------------- Word ----------------
doc = docx.Document(SRC)
body = doc.element.body
for child in list(body):
    if not child.tag.endswith("}sectPr"):
        body.remove(child)
sec = doc.sections[0]
sec.top_margin = Cm(1.2); sec.bottom_margin = Cm(1.2)
sec.left_margin = Cm(1.3); sec.right_margin = Cm(1.3)

def setfont(run, size=10, bold=False, color=None, name="Yu Gothic"):
    run.font.size = Pt(size); run.bold = bold; run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color: run.font.color.rgb = RGBColor.from_string(color)

def para(text="", size=10, bold=False, color=None, space_after=3, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if align is not None: p.alignment = align
    if text: setfont(p.add_run(text), size, bold, color)
    return p

def h1(text, size=11.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    setfont(p.add_run(text), size, True, "1F4E79")
    pPr = p._element.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "8")
    bt.set(qn("w:space"), "1"); bt.set(qn("w:color"), "1F4E79")
    pbdr.append(bt); pPr.append(pbdr)
    return p

def bullet(text, size=9.5, bold=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.5); p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.line_spacing = 1.0
    setfont(p.add_run("・" + text), size, bold, color)
    return p

def borders(tbl, color="808080"):
    b = OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"4")
        e.set(qn("w:space"),"0"); e.set(qn("w:color"),color)
        b.append(e)
    tbl._tbl.tblPr.append(b)

def shade(cell, color):
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"),"clear")
    sh.set(qn("w:fill"), color); cell._tc.get_or_add_tcPr().append(sh)

def table(rows, header=True, size=9, widths=None, headfill="DDEBF7", after=2):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = doc.styles["Normal Table"]
    borders(t)
    for r in t.rows:
        r._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0.5); p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.line_spacing = 1.0
            setfont(p.add_run(str(val)), size, bold=(header and i == 0))
            if header and i == 0: shade(c, headfill)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(after)
    return t

def drop_last_empty():
    e = doc.paragraphs[-1]
    e._element.getparent().remove(e._element)

def pict(path, cm=18.4):
    doc.add_picture(path, width=Cm(cm))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)

# ===== 1ページ目 =====
para("手術室オカレンス・サマリー（2026年9月度）", 15, True, "1F4E79", space_after=0)
para("作成：麻酔科　玉里順子", 9, color="404040", space_after=1, align=WD_ALIGN_PARAGRAPH.RIGHT)
para("集計対象：麻酔科／CE／NS オカレンスレポート 2026年9月分　　業務委員会 配布資料　　作成日：2026年10月13日",
     8.5, color="595959", space_after=4)

ops_txt = f"{ops9} 件" if ops9 else f"{WAIT}"
aops_txt = f"{aops9} 件" if aops9 else WAIT
rate_txt = f"{r100(T9, ops9)} 件" if ops9 else WAIT
prev_rate = r100(T8, OPS[8])
kpi = table([
    ["総手術件数", f"{ops_txt}\n（麻酔科管理 {aops_txt}）", "総オカレンス件数", f"{T9} 件\n（前月 {T8} 件）",
     "100手術あたり報告件数", f"{rate_txt}\n（前月 {prev_rate} 件）"],
], header=False, size=10, widths=[2.4,3.4,3.0,2.8,3.4,3.2], after=2)
for j in (0, 2, 4):
    shade(kpi.cell(0, j), "DDEBF7")
    for r in kpi.cell(0, j).paragraphs:
        for run in r.runs: run.bold = True

a9, c9 = anes[-1], ce[-1]
def rr(n, d, lab): return f"{r100(n, d)}％（÷{lab}{d}）" if d else "―"
table([
    ["職種","麻酔科","CE","NS","合計"],
    ["9月", f"{a9}（前月{anes[-2]}）", f"{c9}（前月{ce[-2]}）", f"{ns_s}（前月{ns[-2]}）", f"{T9}（前月{T8}）"],
    ["報告率", rr(a9, aops9, "麻酔科管理"), rr(c9, ops9, ""), rr(ns_s, ops9, ""), rr(T9, ops9, "")],
], size=9, widths=[1.8,5.0,3.8,3.8,4.0], after=4)

h1("1. 月別推移（2025年4月〜2026年9月）")
pict(trend_png)

h1("2. 手術件数と報告件数（100手術あたり）")
pict(rate_png)

h1("3. 今月のポイント")
if ops9:
    p1 = (f"総数は{T8}→{T9}件（{T9-T8:+d}件）で、上半期の月平均（{h1avg:.0f}件）並み。100手術あたり{r100(T9, ops9)}件"
          f"（前月{prev_rate}件）。3職種とも報告が継続して集まっており、報告文化は定着している。")
else:
    p1 = (f"総数は{T8}→{T9}件（{T9-T8:+d}件）で、上半期の月平均（{h1avg:.0f}件）並み。3職種とも報告が継続して集まっており、"
          f"報告文化は定着している。")
bullet(p1, 9.5, bold=True)
bullet(f"麻酔科{a9}件（前月{anes[-2]}）。「手術準備不足」31件（実報告の{31/a9*100:.1f}％）は、8月に多かったBISモニター（6→2件）・ハサミ（5→1件）が減り、"
       "ボールペン・マジック・蛇管立て・人工鼻（各4件）など文具・小物が上位に。薬剤9件中6件は病棟からのルートトラブル。", 9.5)
bullet(f"CE{c9}件のうち故障・不良は9件（{9/c9*100:.1f}％）で、設定・操作12件＋接続・電源7件の計19件（{19/c9*100:.1f}％）は現場で設定や接続を直して解決している。"
       "気腹装置のモード選択2件・エリートの入力/表示設定6件が目立つ。8月に7回反復したAtlanのフロー計測不良は9月は報告なし。", 9.5)
bullet(f"NS{ns_s}件。経鼻挿管で使用した綿棒が鼻腔内に残ったまま帰室し、患者の違和感の訴えで発見された（9/28）。6/4の鼻内ジョグ紛失に続く鼻腔内の遺残で、"
       "件数と切り離して重点対応する。", 9.5, bold=True)
bullet("9/15 9roomのハイブリッド連動エラー（入室後に手術中止、9roomが1日使用不可）は、CE・NS・麻酔科の3部門から報告が上がり、部門をまたいで共有された。", 9.5)

# ===== 2ページ目 =====
doc.add_page_break()

h1(f"4. 分類別 × 職種別（9月・延べ{NOBE}件）")
pict(cat_png)

h1("5. Good Job / Good Catch（患者に到達する前に止められた事例）")
table([
    ["発生日","部署","内容","評価"],
    ["9/7","NS・眼科","ベガモックスの希釈方法の誤りを、投与前に先輩看護師が指摘し作り直した","Good Catch"],
    ["9月","麻酔科","薬疹の既往がある抗菌薬がオーダーされていることに気付き、報告した","Good Catch"],
    ["9/29","NS・脳神経外科","部屋準備の配置が左右逆であることに患者入室前に気付き、準備し直した","Good Catch"],
    ["9/11","CE・外科","VIO3の通電不良を脂肪組織（高インピーダンス）とSoft凝固の組合せと判断し、Swift凝固へ変更して止血","Good Job"],
    ["9/10","CE・脳神経外科","脳波計のハム雑音に対し、ケーブル類を絶縁体で覆ってノイズを軽減した","Good Job"],
    ["9/2","NS","2年ほど放置されていた未使用のCO2ボンベを見つけて返却した","Good Job"],
], size=8.2, widths=[1.4,2.6,11.2,1.8])
drop_last_empty()

h1("6. 注目事象と今月の重点対策（10月）")
table([
    ["#","注目事象（9月）","重点対策","主担当","期限"],
    ["1","【9/28 口腔外科】経鼻挿管の綿棒が鼻腔内に残ったまま帰室。患者の違和感の訴えで発見、後日抜去","経鼻挿管の綿棒・鼻内の詰め物は本数を記録し、退室前に抜去を確認","麻酔科・手術室NS","10月中"],
    ["2","【9/9 立ち会い業者】受付で患者名を言わず部屋も確認せず、別の患者の部屋に来た。携帯の持ち込みも多い","患者名・部屋の申告と携帯持ち込み禁止を業者へ文書で通知し、受付で確認","手術室NS・整形外科","10月中"],
    ["3","【9/14・9/18 麻酔器】CO2アラームを「OFF」にして解消（前症例で設定変更の可能性）。呼吸検出なしアラームは下限設定が高かった","アラームはOFFにせず原因を確認。変えた設定は症例終了時に初期値へ戻す","CE・麻酔科","10月中"],
    ["4","【9/15 9room】透視装置と手術台の連動エラーで入室後に手術中止、9roomが1日使用不可（定期点検直後）","両メーカー合同の原因報告と、定期点検後の連動確認","CE・放射線科","10月中"],
    ["5","【9月 麻酔科・NS】タイムアウト未実施1件、同意書の取り忘れ・持参忘れ・チェック漏れ3件、術式変更後の同意書なし","サインインで各同意書と術式変更の反映を読み上げ確認。タイムアウトを省略しない","手術室NS・麻酔科・各科","10月中"],
    ["6","【CE】設定・操作による報告12件（気腹装置のモード、エリートの入力ポート、電気メスの凝固モード）","使う前に見る設定を機器に掲示","CE・手術室NS","10月中"],
    ["7","【麻酔科】手術準備不足31件（ボールペン・マジック・蛇管立て・人工鼻 各4件）","上位品目を部屋ごとに定数化し、朝に残数確認","手術室NS・麻酔科","継続"],
], size=8.2, widths=[0.6,7.2,5.6,2.6,1.4])
drop_last_empty()

h1("7. オカレンス定期カンファレンスの開始（10月〜）")
para("収集開始から1年半が経過し、報告文化は安定してきている。報告を受けた担当者は即時に現場へフィードバックするよう努めているが、なかなか改善しないエラーもある。"
     "そこで今月から月1回の定期カンファレンスを開き、対策の確認と現場からの意見の汲み取りを行う。参加者は固定せず自由参加、形式も自由とする。",
     9.0, space_after=2)

para(f"※麻酔科＝総計{ANES_TOTAL[9]}件－オカレンスなし{ANES_NONE[9]}件。分類別は延べ{NOBE}件（麻酔科は原資料で3件以上の項目を分類に対応づけた）。"
     "件数の増加は報告文化の定着として評価し、安全性は事象の重大性と再発の有無で判断する。",
     7.8, color="595959", space_after=0)

doc.save(OUT)
print("saved:", OUT, os.path.getsize(OUT), "延べ", NOBE, "上位2", top2)
