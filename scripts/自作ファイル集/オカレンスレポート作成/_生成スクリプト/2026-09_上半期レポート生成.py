# -*- coding: utf-8 -*-
"""2026年度上半期 手術室実績とオカレンスレポート集計（10月業務委員会）
使い方: python 2026-09_上半期レポート生成.py <出力docx> <図の作業フォルダ>
9月の手術件数は occ_2026_data.py の OPS[9]／ANES_OPS[9] に入れて再実行する。"""
import os, sys
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import rcParams
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SCR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, SCR)
from occ_2026_data import *

rcParams["font.family"] = "Yu Gothic"
rcParams["axes.unicode_minus"] = False
BASE = os.path.dirname(SCR)
SRC = os.path.join(BASE, "手術室オカレンス・マンスリーレポート_2026年6月度配布用.docx")
OUT, WORK = sys.argv[1], sys.argv[2]
WAIT = "入力待ち"
MS = list(range(4, 10))
FULL = OPS[9] is not None and ANES_OPS[9] is not None

# ---------- 集計 ----------
ix = {m: MONTHS.index(f"2026/{m}") for m in MS}
a26 = {m: ANES[ix[m]] for m in MS}; c26 = {m: CE[ix[m]] for m in MS}; n26 = {m: NS[ix[m]] for m in MS}
t26 = {m: a26[m] + c26[m] + n26[m] for m in MS}
A_H1, C_H1, N_H1 = sum(a26.values()), sum(c26.values()), sum(n26.values())
T_H1 = A_H1 + C_H1 + N_H1
ATOT_H1 = sum(ANES_TOTAL.values())
ops5 = sum(OPS[m] for m in range(4, 9)); ops5_25 = sum(OPS_2025[m] for m in range(4, 9))
aop5 = sum(ANES_OPS[m] for m in range(4, 9)); aop5_25 = sum(ANES_OPS_2025[m] for m in range(4, 9))
OPS25, AOPS25 = sum(OPS_2025.values()), sum(ANES_OPS_2025.values())
OPS24, AOPS24 = sum(OPS_2024.values()), sum(ANES_OPS_2024.values())
if FULL:
    OPSH, AOPSH = sum(OPS.values()), sum(ANES_OPS.values())

def pct(a, b): return f"{(a-b)/b*100:+.1f}％"
def r100(n, d): return f"{n/d*100:.1f}" if d else "―"
def v(x): return str(x) if x is not None else WAIT

cat_h1 = {}
for m in MS:
    for k, n in ANES_BY_CAT[m].items(): cat_h1[k] = cat_h1.get(k, 0) + n
NOBE_A = sum(cat_h1.values())
top = sorted(cat_h1.items(), key=lambda x: -x[1])[:10]

# ---------- 図 ----------
def save(fig, name):
    p = os.path.join(WORK, name); fig.tight_layout(); fig.savefig(p); plt.close(fig); return p

# 図1 手術件数 前年比
fig, ax = plt.subplots(figsize=(9.4, 2.05), dpi=220)
xs = range(len(MS)); w = 0.38
b25 = [OPS_2025[m] for m in MS]; b26 = [OPS[m] or 0 for m in MS]
ax.bar([x - w/2 for x in xs], b25, w, color="#bfbfbf", label="2025年度")
ax.bar([x + w/2 for x in xs], b26, w, color="#1f4e79", label="2026年度")
for x, m in zip(xs, MS):
    ax.text(x - w/2, b25[x] + 8, str(b25[x]), ha="center", fontsize=7.5, color="#595959")
    if OPS[m]: ax.text(x + w/2, OPS[m] + 8, str(OPS[m]), ha="center", fontsize=7.5, color="#1f4e79", fontweight="bold")
    else: ax.text(x + w/2, 30, WAIT, ha="center", fontsize=7.5, color="#c00000", rotation=90)
ax.set_xticks(list(xs)); ax.set_xticklabels([f"{m}月" for m in MS], fontsize=8.5)
ax.set_ylabel("総手術件数", fontsize=9); ax.set_ylim(0, 720); ax.tick_params(axis="y", labelsize=8)
ax.legend(fontsize=8.5, ncol=2, loc="upper left", frameon=False)
for s in ("top", "right"): ax.spines[s].set_visible(False)
ax.grid(axis="y", ls=":", color="#cccccc")
fig_ops = save(fig, "h1_ops.png")

# 図2 職種別推移
tot = [(a or 0)+(c or 0)+(n or 0) for a, c, n in zip(ANES, CE, NS)]
fig, ax = plt.subplots(figsize=(9.4, 2.25), dpi=220)
x = list(range(len(MONTHS)))
ax.axvspan(11.5, len(MONTHS)-0.5, color="#fff2cc", alpha=0.6, lw=0)
ax.text(14.5, max(tot)+22, "2026年度上半期", ha="center", fontsize=8.5, color="#7f6000")
ax.bar(x, tot, color="#e8eef7", label="合計", width=0.62)
ax.plot(x, ANES, marker="o", ms=4, lw=1.8, color="#1f4e79", label="麻酔科")
ax.plot(x, CE, marker="s", ms=4, lw=1.8, color="#c00000", label="CE")
ax.plot(x, [n if n is not None else float("nan") for n in NS], marker="^", ms=4, lw=1.8, color="#2e7d32", label="NS")
ax.set_xticks(x); ax.set_xticklabels(MONTHS, rotation=45, fontsize=8)
ax.set_ylabel("報告件数", fontsize=9); ax.tick_params(axis="y", labelsize=8)
ax.set_ylim(0, max(tot)+40); ax.grid(axis="y", ls=":", color="#cccccc")
ax.legend(fontsize=8.5, ncol=4, loc="upper left", frameon=False)
for s in ("top", "right"): ax.spines[s].set_visible(False)
fig_trend = save(fig, "h1_trend.png")

# 図3 麻酔科 分類別上位10（前年同期比）
fig, ax = plt.subplots(figsize=(9.4, 2.6), dpi=220)
labs = [f"{k}. {ANES_CAT[k]}" for k, _ in top][::-1]
v26 = [n for _, n in top][::-1]; v25 = [ANES_BY_CAT_2025_H1.get(k, 0) for k, _ in top][::-1]
y = range(len(labs)); h = 0.38
ax.barh([i + h/2 for i in y], v26, h, color="#1f4e79", label="2026年度上半期")
ax.barh([i - h/2 for i in y], v25, h, color="#bfbfbf", label="2025年度上半期")
for i in y:
    ax.text(v26[i] + 1.5, i + h/2, str(v26[i]), va="center", fontsize=7.5, color="#1f4e79", fontweight="bold")
    ax.text(v25[i] + 1.5, i - h/2, str(v25[i]), va="center", fontsize=7.5, color="#595959")
ax.set_yticks(list(y)); ax.set_yticklabels(labs, fontsize=8)
ax.set_xlim(0, 205); ax.tick_params(axis="x", labelsize=8); ax.set_xlabel("件数", fontsize=8.5)
ax.legend(fontsize=8.5, loc="lower right", frameon=False)
for s in ("top", "right"): ax.spines[s].set_visible(False)
fig_cat = save(fig, "h1_cat.png")

# 図4 手術準備不足の月別推移
fig, ax = plt.subplots(figsize=(9.4, 1.7), dpi=220)
lab = [f"{m}月" for m in MS]
ax.plot(lab, PREP_2025, marker="o", lw=1.8, color="#a6a6a6", label="2025年度")
ax.plot(lab, PREP_2026, marker="o", lw=2.2, color="#c00000", label="2026年度")
for i, (p, q) in enumerate(zip(PREP_2025, PREP_2026)):
    ax.text(i, q + 2.5, str(q), ha="center", fontsize=8, color="#c00000", fontweight="bold")
    ax.text(i, p - 6.5, str(p), ha="center", fontsize=7.5, color="#7f7f7f")
ax.set_ylim(0, 60); ax.set_ylabel("件数", fontsize=9); ax.tick_params(labelsize=8)
ax.legend(fontsize=8.5, ncol=2, loc="upper left", frameon=False)
ax.grid(axis="y", ls=":", color="#cccccc")
for s in ("top", "right"): ax.spines[s].set_visible(False)
fig_prep = save(fig, "h1_prep.png")

# ---------------- Word ----------------
doc = docx.Document(SRC)
body = doc.element.body
for child in list(body):
    if not child.tag.endswith("}sectPr"): body.remove(child)
sec = doc.sections[0]
sec.top_margin = Cm(1.5); sec.bottom_margin = Cm(1.5)
sec.left_margin = Cm(1.6); sec.right_margin = Cm(1.6)

def setfont(run, size=10, bold=False, color=None, name="Yu Gothic"):
    run.font.size = Pt(size); run.bold = bold; run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color: run.font.color.rgb = RGBColor.from_string(color)

def para(text="", size=10, bold=False, color=None, space_after=3, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    if align is not None: p.alignment = align
    if text: setfont(p.add_run(text), size, bold, color)
    return p

def h1(text, size=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    setfont(p.add_run(text), size, True, "1F4E79")
    pPr = p._element.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    for k, val in (("w:val", "single"), ("w:sz", "8"), ("w:space", "1"), ("w:color", "1F4E79")): bt.set(qn(k), val)
    pbdr.append(bt); pPr.append(pbdr)
    return p

def h2(text):
    p = para(text, 10.5, True, "1F4E79", space_after=2)
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.keep_with_next = True
    return p

def bullet(text, size=10, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.left_indent = Cm(0.45); p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.line_spacing = 1.05
    setfont(p.add_run("・" + text), size, bold)
    return p

def borders(tbl, color="808080"):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        for k, val in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", color)): e.set(qn(k), val)
        b.append(e)
    tbl._tbl.tblPr.append(b)

def shade(cell, color):
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(sh)

def table(rows, size=9, widths=None, after=3, bold_last=False, center_from=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = doc.styles["Normal Table"]; borders(t)
    for r in t.rows: r._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0.5); p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.line_spacing = 1.0
            if center_from is not None and j >= center_from: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            last = bold_last and i == len(rows) - 1
            setfont(p.add_run(str(val)), size, bold=(i == 0 or last),
                    color=("C00000" if str(val) == WAIT else None))
            if i == 0: shade(c, "DDEBF7")
            elif last: shade(c, "F2F2F2")
    if widths:
        for j, wd in enumerate(widths):
            for row in t.rows: row.cells[j].width = Cm(wd)
    doc.add_paragraph().paragraph_format.space_after = Pt(after)
    return t

def pict(path, cm=17.6):
    doc.add_picture(path, width=Cm(cm))
    p = doc.paragraphs[-1]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)

# ===== 表題 =====
para("2026年度上半期　手術室実績とオカレンスレポート集計", 16, True, "1F4E79", space_after=0)
para("作成：麻酔科　玉里順子", 9.5, color="404040", space_after=1, align=WD_ALIGN_PARAGRAPH.RIGHT)
para("対象期間：2026年4月〜9月　　業務委員会（10月）資料　　作成日：2026年10月13日", 8.5, color="595959", space_after=4)

# ===== ① 手術実績 =====
h1("① 2026年度上半期 手術実績")
rows = [["月", "総手術件数", "前年同月", "前年比", "麻酔科管理", "前年同月", "前年比"]]
for m in MS:
    rows.append([f"{m}月", v(OPS[m]), OPS_2025[m], pct(OPS[m], OPS_2025[m]) if OPS[m] else "―",
                 v(ANES_OPS[m]), ANES_OPS_2025[m], pct(ANES_OPS[m], ANES_OPS_2025[m]) if ANES_OPS[m] else "―"])
if FULL:
    rows.append(["計", OPSH, OPS25, pct(OPSH, OPS25), AOPSH, AOPS25, pct(AOPSH, AOPS25)])
else:
    rows.append(["4〜8月計", ops5, ops5_25, pct(ops5, ops5_25), aop5, aop5_25, pct(aop5, aop5_25)])
    rows.append(["上半期計", WAIT, OPS25, "―", WAIT, AOPS25, "―"])
table(rows, size=9.5, widths=[2.4, 2.7, 2.4, 2.2, 2.7, 2.4, 2.2], bold_last=True, center_from=1)
para(f"（参考）2024年度上半期：総手術件数 {OPS24}件／麻酔科管理 {AOPS24}件", 8.5, color="595959", space_after=3)
pict(fig_ops)
if FULL:
    bullet(f"上半期の総手術件数は{OPSH}件で前年同期（{OPS25}件）から{OPSH-OPS25:+d}件（{pct(OPSH, OPS25)}）、"
           f"麻酔科管理は{AOPSH}件で{AOPSH-AOPS25:+d}件（{pct(AOPSH, AOPS25)}）。総手術件数は6か月すべてで前年同月を上回った。")
else:
    bullet(f"4〜8月の総手術件数は{ops5}件で前年同期（{ops5_25}件）から{ops5-ops5_25:+d}件（{pct(ops5, ops5_25)}）、"
           f"麻酔科管理は{aop5}件で{aop5-aop5_25:+d}件（{pct(aop5, aop5_25)}）。全ての月で前年を上回った。")
bullet("7月の617件（麻酔科管理388件）は上半期で最も多く、前年同月比＋21.9％。")

# ===== ② オカレンス集計 =====
h1("② 上半期オカレンスレポート集計")
rows = [["区分", "2026年度上半期", "2025年度上半期", "増減"],
        ["麻酔科（レポート総数）", f"{ATOT_H1}件", f"{ANES_TOTAL_2025_H1}件", f"{ATOT_H1-ANES_TOTAL_2025_H1:+d}件"],
        ["　うちオカレンス有", f"{A_H1}件（{A_H1/ATOT_H1*100:.1f}％）", f"{ANES_REAL_2025_H1}件（{ANES_REAL_2025_H1/ANES_TOTAL_2025_H1*100:.1f}％）", f"{A_H1-ANES_REAL_2025_H1:+d}件"],
        ["CE", f"{C_H1}件", f"{CE_2025_H1}件", f"{C_H1-CE_2025_H1:+d}件"],
        ["NS", f"{N_H1}件（4〜9月）", f"{NS_2025_H1}件（6〜9月）", "―"],
        ["3職種合計", f"{T_H1}件", "―", "―"]]
table(rows, size=9.5, widths=[4.6, 4.6, 4.6, 3.0], bold_last=True)

h2("● 月別・職種別の報告件数と100手術あたり報告件数")
rows = [["月", "麻酔科", "CE", "NS", "合計", "総手術件数", "100手術あたり"]]
for m in MS:
    rows.append([f"{m}月", a26[m], c26[m], n26[m], t26[m], v(OPS[m]), r100(t26[m], OPS[m]) if OPS[m] else "―"])
rows.append(["上半期計", A_H1, C_H1, N_H1, T_H1, OPSH if FULL else WAIT, r100(T_H1, OPSH) if FULL else "―"])
table(rows, size=9.5, widths=[2.4, 2.2, 2.2, 2.2, 2.2, 3.0, 3.0], bold_last=True, center_from=1)
pict(fig_trend)
bullet(f"3職種で上半期{T_H1}件、月平均{T_H1/6:.0f}件。最多は8月（{t26[8]}件）、最少は5月（{t26[5]}件）。"
       "6月度から3職種のレポートを統合した月報を業務委員会に提出し、毎月、重点対策を担当・期限つきで示した。", bold=True)
bullet(f"麻酔科は{A_H1}件で前年同期（{ANES_REAL_2025_H1}件）とほぼ同数。ただし麻酔科管理件数が前年より増えているため（"
       + (f"{AOPSH}件、{pct(AOPSH, AOPS25)}" if FULL else "4〜8月で約13％増") + "）、"
       + (f"麻酔科管理100件あたりでは{ANES_REAL_2025_H1/AOPS25*100:.1f}→{A_H1/AOPSH*100:.1f}件と下がっている。" if FULL else "1症例あたりの報告頻度は下がっている。") + "レポート総数（オカレンスなしを含む）は前年より増えており、提出そのものは続いている。")
bullet(f"CEは{C_H1}件で前年同期（{CE_2025_H1}件）より{CE_2025_H1-C_H1}件少ない。4〜6月は月23〜30件、7〜9月は月32〜38件と、"
       "年度後半に向けて報告が戻っている。")
bullet(f"NSは5月に8件まで減ったが、7〜9月は月18〜23件に回復した。2025年度上半期（6〜9月、月平均{NS_2025_H1/4:.0f}件）より少なく、"
       "報告件数の減少はまず過少報告を疑う。現場が拾い上げた事象を書き残せる時間と様式の確保が課題。")

# ===== ③ 麻酔科 分類別 =====
h1("③ 麻酔科オカレンスの分類別（上位10分類・前年同期比）")
pict(fig_cat)
p34, p28 = cat_h1.get(34, 0), cat_h1.get(28, 0)
bullet(f"上位2分類（34.手術準備不足{p34}件＋28.ME機器・ディスポ不良{p28}件）で実件数{A_H1}件の{(p34+p28)/A_H1*100:.1f}％。"
       f"前年同期（{ANES_BY_CAT_2025_H1[34]}＋{ANES_BY_CAT_2025_H1[28]}件、実件数の{(ANES_BY_CAT_2025_H1[34]+ANES_BY_CAT_2025_H1[28])/ANES_REAL_2025_H1*100:.1f}％）と同じ構造が続いている。")
bullet(f"前年より減ったのは19.手術時間の延長（{ANES_BY_CAT_2025_H1[19]}→{cat_h1.get(19,0)}件）と43.その他（{ANES_BY_CAT_2025_H1[43]}→{cat_h1.get(43,0)}件）。"
       f"増えたのは3.入退室（{ANES_BY_CAT_2025_H1[3]}→{cat_h1.get(3,0)}件）と39.義歯・指輪等の取り忘れ（{ANES_BY_CAT_2025_H1[39]}→{cat_h1.get(39,0)}件）。")

h2("● 34.手術準備不足の月別推移")
pict(fig_prep)
bullet("2025年度は麻酔カートの見直しにより、7月49件→9月19件と減った。2026年度は4〜5月から24件と多く、6月43件を山に、7〜9月は27〜35件で推移している。")
bullet("欠品の品目は、9月にBISモニター（8月6件→2件）・ハサミ（5件→1件）が減った。一方、ボールペン・マジックは7月から毎月上位に残り、9月は蛇管立て・人工鼻（各4件）も並んだ。"
       "品目が固定してきており、品目を絞った定数化が効く段階にある。")

# ===== ④ CE・NS =====
h1("④ CE・NSオカレンスの特徴")
h2("● CE（上半期183件）")
kind = ce_kind()
bullet("4〜6月は月23〜30件と前年より少なかった。7月は接続・電源の外れ（8件）と術中の機器不良、8月は麻酔器12件（うちAtlanのフロー計測不良・リークテスト不合格が7回、"
       "回路内の水滴が共通要因）と日本光電モニタ6件に集中した。")
bullet(f"9月は32件のうち故障・不良が{kind['故障・不良']}件で、設定・操作{kind['設定・操作']}件と接続・電源{kind['接続・電源']}件（計{kind['設定・操作']+kind['接続・電源']}件）は"
       "現場で設定や接続を直して解決している。気腹装置の成人/小児モード、エリートの入力ポート、電気メスの凝固モードなど、使う前に見る設定が繰り返し報告されている。")
bullet("8月に7回反復したAtlanのフロー計測不良は9月は報告がなく、7/27・8/3と続いた10roomのLAN・無停電装置の障害も9月は再発していない。")

h2(f"● NS（上半期{N_H1}件）")
nsd = {}
for r in NS_ROWS: nsd[r[6]] = nsd.get(r[6], 0) + 1
rows = [["分類"] + GROUPS + ["計"], ["件数"] + [nsd.get(g, 0) for g in GROUPS] + [len(NS_ROWS)]]
NS_NOTE = f"※内容を確認できた{len(NS_ROWS)}件を分類（5月分のうち{N_H1-len(NS_ROWS)}件は内訳なし）。"
table(rows, size=9, widths=[1.6, 3.2, 1.4, 1.4, 3.0, 2.2, 1.6, 1.4], center_from=1)
para(NS_NOTE, 8.5, color="595959", space_after=2)
bullet("確認不足・準備不足が37件（37％）で最も多い。病棟⇔手術室、診療科⇔手術室の事前連絡に関するもの（術式変更・CT追加・借り物器械・デモ機器・PM/ICD症例など）が"
       "12件あり、4月から8月まで毎月続いた。")
bullet("職員間の言動に関する報告（術中の威圧的な叱責、業務と無関係な撮影・容姿への発言、不適切な発言と身体への接触、抜管前の帰室指示に「言い合える関係性がない」）が"
       "5月・7月・8月に計4件あった。報告として上がったこと自体が重要で、内容は個別に対応する。")

# ===== ⑤ 上半期に繰り返した事象 =====
h1("⑤ 上半期に繰り返した事象（3職種の横断）")
table([
    ["事象", "上半期の報告", "主な内容"],
    ["体内遺残・カウント不一致", "NS 4件\n麻酔科 33.遺残等 2件", "6/4 鼻内のジョグが術後に見つからない、6/30 ガーゼカウント不一致、8/25 ラインガーゼのカウント不一致、9/28 経鼻挿管の綿棒が鼻腔内に残ったまま帰室"],
    ["手術部位・患者・術式の確認", "NS 6件", "4/8 オーダー入力時の患者間違い、6/5 器械が他患者分に紛れる、8/7 反対側の目を消毒、8/19 抜歯本数の不足、9/9 業者が別の患者の部屋へ、9/29 部屋配置の左右逆"],
    ["薬剤", "NS 6件\n麻酔科 22.薬剤等 49件", "6/9 アセリオの重複投与、6/30 骨髄用ゲンタシンの静注、7/27 アイリーアの投与量誤り（いずれも患者到達）、9/7 ベガモックス希釈の誤り（投与前に発見）"],
    ["麻薬の管理", "NS 2件・麻酔科 3件", "8/3 使用後の麻薬を金庫に返却せず一晩保管、8/4 不足に気付きながら未確認で帰宅、9月 麻薬持参忘れ2件・準備忘れ1件"],
    ["針刺し", "NS 3件\n麻酔科 41.針刺し 4件", "8/31 研修医が2週続けて縫合針で針刺し、8/31 カテ室の片付け時、9/10 新品の糸針セットから飛び出した針"],
    ["転落・ずり落ち", "NS 2件・麻酔科 1件", "4/9 カテ室でストレッチャーからずり落ち、5/15 マイトラクリップ抜管直後にベッドから転落、9月 ダヴィンチのアーム接触後にずり落ち"],
    ["設備・ネットワーク", "NS・CE", "7/27・8/3 10roomのLAN障害（無停電装置）、9/2 天井の水漏れ、9/15 9roomのハイブリッド連動エラー、9/28 10roomの空調故障"],
], size=8.6, widths=[3.6, 3.2, 10.4])

# ===== ⑥ 対策の状況と下半期の重点 =====
h1("⑥ 上半期に示した重点対策の状況と、下半期の重点")
table([
    ["重点対策（提示月）", "9月時点の状況"],
    ["麻酔準備の欠品の定数化（7月〜）", "7月27件→8月35件→9月31件。BIS・ハサミは減少、ボールペン・マジックは上位に残る"],
    ["麻酔器（Atlan）の水滴対策（8月）", "8月7回→9月0件"],
    ["10roomのLAN・無停電装置（7・8月）", "9月は再発なし"],
    ["手術部位（左右）の確認（8月）", "9月に部屋配置の左右逆1件（患者入室前に修正）"],
    ["麻薬管理（8月）", "NSからの再発報告なし。麻酔科で麻薬の持参忘れ2件・準備忘れ1件"],
    ["針刺し防止（8月）", "9月にも1件（新品の糸針セット開封時）"],
], size=9, widths=[6.6, 10.6])

h2("● 下半期の重点")
bullet("体内遺残の防止：ガーゼ・器械に加え、鼻内の詰め物・経鼻挿管の綿棒もカウントの対象にし、退室前に抜去を確認する。", bold=True)
bullet("手術部位・患者の確認：消毒直前の左右の読み合わせ、サインアウトでの本数・部位の突合、立ち会い業者の受付での患者名・部屋の申告。")
bullet("報告の回復：NS・CEは前年同期より少ない。報告しやすい様式と時間を確保し、出された報告には月報で対応を返す。")
bullet("麻酔準備の欠品：上位品目を部屋ごとに定数化し、朝の残数確認に組み込む。月ごとの件数で効果を見る。")
bullet("機器の設定確認とアラーム：使う前に見る設定を機器に掲示し、アラームはOFFにせず原因を確認する。")
bullet("事前連絡と職員間の関係：術式変更・借り物器械・デモ機器・PM/ICD症例の連絡先と期限を各科へ周知する。疑問を口に出せる場（症例前ブリーフィング）を続ける。")

h1("⑦ 今後の取り組み：オカレンス定期カンファレンスの開始")
para("オカレンスレポートの収集開始から1年半が経過し、報告文化は安定してきている。報告を受けた担当者は、即時に現場へフィードバックするよう努めているが、"
     "なかなか改善しないエラーもある。そこで今月から月1回の定期カンファレンスを開き、対策の進み具合を確認するとともに、現場からさらに意見を汲み取る。", 10, space_after=3)
table([
    ["項目", "内容"],
    ["開始", "2026年10月から"],
    ["頻度", "月1回"],
    ["目的", "改善しにくいエラーの確認と、現場からの意見の汲み取り"],
    ["参加者", "固定せず、自由参加"],
    ["形式", "自由"],
], size=9.5, widths=[3.0, 14.2])
para("※麻酔科の件数は「レポート総数－99.オカレンスなし」。手術件数は医事課提供値。オカレンスレポートは報告件数の増加を報告文化の定着として評価し、"
     "件数の減少はまず過少報告を疑う。安全性は個々の事象の重大性と再発の有無で判断する。", 8, color="595959", space_after=0)

doc.save(OUT)
print("saved", OUT, "FULL" if FULL else "PARTIAL", "T_H1", T_H1)
