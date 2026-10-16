# -*- coding: utf-8 -*-
"""Word の表を「短い項目は 1 行」に収める列幅・文字サイズの自動調整。

使い方:  python fit_tables.py 入力.docx [出力.docx] [--pdf]
  出力を省くと上書き。--pdf で同名 PDF も作る（Word COM・SaveAs2）。

手順
  1. 各セルの 1 行幅をフォントで実測する（PIL、メイリオ/游ゴシック等）。
  2. 表全体の幅は変えずに列幅を配り直す。評価は「表の高さ（行数）＋短い項目が折れた罰点」。
     見出し行・左端の項目名（表幅の 45% まで）・表幅の 1/4 以下の項目は折れると重い罰点、
     1/4〜0.55 は軽い罰点。
     列の間で幅をやり取りする局所探索で、この値が最小になる配分を選ぶ。
  3. 文字を 5% ずつ下げた案（元の 0.9 倍まで）も比べ、折れが減るときだけ採る。
  4. Word で実際の行数を数え、まだ折れている短い項目は列を少しずつ広げる。
"""
import sys, os, re, math, shutil, tempfile
from docx import Document
from docx.shared import Pt, Emu
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import ImageFont

FONTDIR = r"C:\Windows\Fonts"
FONTFILES = {  # (regular, bold)
    "meiryo": ("meiryo.ttc", "meiryob.ttc"),
    "yu gothic": ("YuGothR.ttc", "YuGothB.ttc"),
    "游ゴシック": ("YuGothR.ttc", "YuGothB.ttc"),
    "ms gothic": ("msgothic.ttc", "msgothic.ttc"),
    "ｍｓ ゴシック": ("msgothic.ttc", "msgothic.ttc"),
    "ms mincho": ("msmincho.ttc", "msmincho.ttc"),
    "ｍｓ 明朝": ("msmincho.ttc", "msmincho.ttc"),
    "游明朝": ("yumin.ttf", "yumindb.ttf"),
}
SAFETY = 1.06          # 推定幅の余裕（Word の実描画との差を吸収）
MIN_SCALE = 0.9        # 文字の縮小は元の 0.9 倍まで（読める大きさを守る）
MIN_PT = 8.5           # どの表もこれより小さくは縮めない
PT_PER_EMU = 1 / 12700
_cache = {}


def font_obj(name, bold, size):
    key = (name, bold, round(size * 4))
    if key not in _cache:
        files = FONTFILES.get((name or "meiryo").lower(), FONTFILES["meiryo"])
        path = os.path.join(FONTDIR, files[1 if bold else 0])
        if not os.path.exists(path):
            path = os.path.join(FONTDIR, FONTFILES["meiryo"][1 if bold else 0])
        _cache[key] = ImageFont.truetype(path, size=max(1, int(round(size * 20))))  # 20 倍で測る
    return _cache[key]


def style_font(doc):
    st = doc.styles["Normal"]
    name = None
    rpr = st.element.rPr
    if rpr is not None and rpr.rFonts is not None:
        name = rpr.rFonts.get(qn("w:eastAsia")) or rpr.rFonts.get(qn("w:ascii"))
    size = st.font.size.pt if st.font.size else 10.5
    return name or "meiryo", size


def run_size(r, default):
    return r.font.size.pt if r.font.size else default


def para_width_pt(p, dfont, dsize, scale=1.0):
    """段落を 1 行で書いたときの幅（pt）。"""
    w = 0.0
    for r in p.runs:
        t = r.text
        if not t:
            continue
        size = run_size(r, dsize) * scale
        if r.font.superscript or r.font.subscript:
            size *= 0.65
        f = font_obj(r.font.name or dfont, bool(r.bold), size)
        w += f.getlength(t) / 20.0
    return w


def cell_margins_pt(tbl):
    """左右のセル余白（pt）。既定 0.19cm。"""
    l = r = 5.4
    mar = tbl._tbl.tblPr.find(qn("w:tblCellMar"))
    if mar is not None:
        for side, nm in (("left", "l"), ("start", "l"), ("right", "r"), ("end", "r")):
            e = mar.find(qn("w:" + side))
            if e is not None and e.get(qn("w:type")) == "dxa":
                v = int(e.get(qn("w:w"))) / 20
                if nm == "l": l = v
                else: r = v
    return l + r + 1.0  # 罫線分


def grid_widths(tbl):
    return [int(gc.get(qn("w:w"))) / 20 for gc in tbl._tbl.tblGrid.findall(qn("w:gridCol"))]


def table_cells(tbl):
    """(行, 列, tc) を列番号つきで列挙。横結合セルは列番号 -1。"""
    out = []
    for ri, tr in enumerate(tbl._tbl.tr_lst):
        ci = 0
        for tc in tr.tc_lst:
            span = tc.tcPr.gridSpan.val if (tc.tcPr is not None and tc.tcPr.gridSpan is not None) else 1
            out.append((ri, ci if span == 1 else -1, tc))
            ci += span
    return out


def measure(tbl, dfont, dsize, scale):
    """[(行, 列, 1 行で書いた本文幅pt)] と左右余白pt。"""
    from docx.table import _Cell
    out = []
    for ri, ci, tc in table_cells(tbl):
        if ci < 0:
            continue
        cell = _Cell(tc, tbl)
        w = max((para_width_pt(p, dfont, dsize, scale) for p in cell.paragraphs), default=0)
        out.append((ri, ci, w * SAFETY))
    return out, cell_margins_pt(tbl)


def lines_of(nat, avail):
    if nat <= 0:
        return 1
    return max(1, math.ceil(nat / max(avail, 1)))


def cost(cells, pad, widths, W):
    """表の高さ（行数）＋短い項目が折れた罰点。小数部は文字の詰まり具合。"""
    rows = {}
    pen = 0.0
    fill = 0.0
    for ri, ci, nat in cells:
        avail = widths[ci] - pad
        n = lines_of(nat, avail)
        rows[ri] = max(rows.get(ri, 1), n)
        if n > 1:
            # 短い項目ほど折れたら重い。見出し行は必ず 1 行
            if ri == 0 or nat <= W * 0.25 or (ci == 0 and nat <= W * 0.45 - pad):
                pen += 6
            elif nat <= W * 0.4:
                pen += 2
            elif nat <= W * 0.55:
                pen += 0.5
        fill += nat / max(avail, 1)
    return sum(rows.values()) + pen + 0.002 * fill


def optimize(cells, pad, ncol, W):
    minw = pad + 22  # 2 文字分
    starts = []
    full = [max([nat for r, c, nat in cells if c == j] or [0]) + pad for j in range(ncol)]
    vol = [sum(nat for r, c, nat in cells if c == j) + pad for j in range(ncol)]
    for base in (full, vol, [1] * ncol, [math.sqrt(v) for v in vol]):
        t = sum(base)
        starts.append([max(minw, W * x / t) for x in base])
    best, bestc = None, None
    for w in starts:
        k = W / sum(w); w = [x * k for x in w]
        c = cost(cells, pad, w, W)
        for step in (85, 42, 21, 10, 5, 2.5):
            improved = True
            while improved:
                improved = False
                cand = None
                for i in range(ncol):
                    for j in range(ncol):
                        if i == j or w[j] - step < minw:
                            continue
                        w2 = w[:]; w2[i] += step; w2[j] -= step
                        c2 = cost(cells, pad, w2, W)
                        if c2 < c - 1e-9 and (cand is None or c2 < cand[0]):
                            cand = (c2, w2)
                if cand:
                    c, w = cand; improved = True
        if bestc is None or c < bestc:
            best, bestc = w, c
    return best, bestc


def set_widths(tbl, widths):
    tblPr = tbl._tbl.tblPr
    lay = tblPr.find(qn("w:tblLayout"))
    if lay is None:
        lay = OxmlElement("w:tblLayout"); tblPr.append(lay)
    lay.set(qn("w:type"), "fixed")
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(int(round(sum(widths) * 20))))
    for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(round(w * 20))))
    for ri, ci, tc in table_cells(tbl):
        if ci < 0:
            continue
        tcPr = tc.get_or_add_tcPr()
        tcW = tcPr.find(qn("w:tcW"))
        if tcW is None:
            tcW = OxmlElement("w:tcW"); tcPr.insert(0, tcW)
        tcW.set(qn("w:type"), "dxa"); tcW.set(qn("w:w"), str(int(round(widths[ci] * 20))))


def scale_fonts(tbl, dsize, scale):
    from docx.table import _Cell
    for ri, ci, tc in table_cells(tbl):
        for p in _Cell(tc, tbl).paragraphs:
            for r in p.runs:
                base = r.font.size.pt if r.font.size else dsize
                r.font.size = Pt(round(base * scale * 2) / 2)


def fit_docx(src, dst):
    doc = Document(src)
    dfont, dsize = style_font(doc)
    sec = doc.sections[0]
    textw = (sec.page_width - sec.left_margin - sec.right_margin) * PT_PER_EMU
    report = []
    intended = set()  # 1 行に収めるはずのセル (表, 行, 列)
    for ti, tbl in enumerate(doc.tables):
        gw = grid_widths(tbl)
        W = min(sum(gw), textw) if sum(gw) > 0 else textw
        ncol = len(gw)
        choice = None
        scale = 1.0
        from docx.table import _Cell as _C
        sizes = [r.font.size.pt if r.font.size else dsize
                 for _ri, _ci, _tc in table_cells(tbl) for _p in _C(_tc, tbl).paragraphs for r in _p.runs if r.text.strip()]
        tmin = min(sizes) if sizes else dsize
        while scale >= MIN_SCALE - 1e-6 and (scale == 1.0 or tmin * scale >= MIN_PT - 1e-6):
            cells, pad = measure(tbl, dfont, dsize, scale)
            w, c = optimize(cells, pad, ncol, W)
            c += (1 - scale) * 40  # 5% 縮めるごとに 2 点。短い項目の折れを消すときだけ選ぶ
            if choice is None or c < choice[1] - 0.5:
                choice = (w, c, scale)
            scale = round(scale - 0.05, 2)
        widths, _, scale = choice
        if scale < 1.0:
            scale_fonts(tbl, dsize, scale)
        set_widths(tbl, widths)
        cells, pad = measure(tbl, dfont, dsize, 1.0)  # 縮小後の run サイズで
        for ri, ci, nat in cells:
            if lines_of(nat, widths[ci] - pad) == 1:
                intended.add((ti + 1, ri + 1, ci + 1))
        report.append((ti + 1, [round(w / 28.35, 2) for w in widths], scale))
    doc.save(dst)
    return report, intended


def word_verify(path, intended, pdf=False, rounds=8):
    """Word で行数を数え、折れた短い項目の列を広げる。"""
    import win32com.client, pythoncom
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False; word.DisplayAlerts = 0
    left = []
    try:
        d = word.Documents.Open(os.path.abspath(path))
        d.Repaginate()
        for ti in range(1, d.Tables.Count + 1):
            t = d.Tables(ti)
            if not t.Uniform:
                continue
            ncol = t.Columns.Count
            for _ in range(rounds):
                bad = {}
                widths = [t.Columns(j).Width for j in range(1, ncol + 1)]
                W = sum(widths)
                for c in t.Range.Cells:
                    txt = c.Range.Text[:-2]
                    if not txt.strip():
                        continue
                    if (ti, c.RowIndex, c.ColumnIndex) in intended and c.Range.ComputeStatistics(1) > 1:
                        bad.setdefault(c.ColumnIndex, []).append(txt)
                if not bad:
                    break
                # 折れた列へ 3pt ずつ。最も広い（折れていない）列から差し引く
                donors = sorted([j for j in range(1, ncol + 1) if j not in bad],
                                key=lambda j: -widths[j - 1])
                if not donors:
                    break
                for j in bad:
                    dj = donors[0]
                    if t.Columns(dj).Width - 3 < 40:
                        break
                    t.Columns(j).Width = t.Columns(j).Width + 3
                    t.Columns(dj).Width = t.Columns(dj).Width - 3
                d.Repaginate()
            else:
                pass
            for c in t.Range.Cells:
                txt = c.Range.Text[:-2]
                if (ti, c.RowIndex, c.ColumnIndex) in intended and c.Range.ComputeStatistics(1) > 1:
                    left.append((ti, c.RowIndex, c.ColumnIndex, txt))
        d.Save()
        if pdf:
            tmp = os.path.join(tempfile.gettempdir(), "fit_tables_out.pdf")
            d.SaveAs2(tmp, FileFormat=17)
            d.Close(False)
            shutil.copyfile(tmp, os.path.splitext(path)[0] + ".pdf")
        else:
            d.Close(False)
    finally:
        word.Quit()
    return left


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    src = args[0]; dst = args[1] if len(args) > 1 else src
    rep, intended = fit_docx(src, dst)
    for ti, w, sc in rep:
        print(f"表{ti}: 列幅(cm) {w}" + (f"  文字 {int(sc*100)}%" if sc < 1 else ""))
    left = word_verify(dst, intended, pdf="--pdf" in sys.argv)
    print("折れたままの短い項目:", len(left))
    for x in left:
        print("  表%d 行%d 列%d: %s" % x)
