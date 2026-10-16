# -*- coding: utf-8 -*-
"""教材 docx を「読める大きさ・一つの書体・崩れない表」に整える後処理。

使い方:  python readable_docx.py 文書.docx [文書2.docx ...] [--no-pdf]
  各文書の隣の _旧版 に「<名前>_見やすさ調整前.docx/.pdf」を控えてから上書きする。

やること
  1. 書体：全スタイル・全ランの書体をメイリオに。テーマ書体（asciiTheme 等）を外す。
     記号用の書体（Symbol・Wingdings・Cambria Math）はそのまま。
  2. 文字の大きさ（上付き・下付きは触らない）
     ・本文：最も多い大きさが 10pt 未満なら 10.5pt になるよう全体を拡大。下限 10pt。
     ・表：4 列以下は 9.5pt、5 列以上は 9pt を下限に。
     ・参考文献（「参考文献」「文献」見出しより後）：下限 9pt。
  3. 図：幅 13.5cm 未満の図を 16cm に（版面の高さを超えない範囲で）。
  4. 表：fit_tables.py で列幅を配り直し、行の頁割れ禁止・見出し行の繰り返し・
     見出し行を次の行と一緒に組む設定を入れる。
  5. Word で目次などのフィールドを更新し、保存して PDF を作り直す。
"""
import sys, os, re, shutil, tempfile, collections
from docx import Document
from docx.shared import Pt, Emu
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import fit_tables as FT

FONT = "Meiryo"
KEEP_FONTS = ("symbol", "wingdings", "cambria math", "webdings", "segoe ui symbol")
BODY_MIN, BODY_TARGET = 10.0, 10.5
REF_MIN = 9.0
FIG_MIN_CM, FIG_CM = 13.5, 16.0
REF_HEAD = re.compile(r"^(参考文献|文献|引用文献|References|参考資料)$")


def default_size(doc):
    try:
        rpr = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
        return int(rpr.find(qn("w:sz")).get(qn("w:val"))) / 2
    except Exception:
        return 10.5


def style_size(st, dflt):
    while st is not None:
        if st.font.size:
            return st.font.size.pt
        st = st.base_style
    return dflt


def eff_size(run, para, dflt):
    if run.font.size:
        return run.font.size.pt
    try:
        if run.style is not None and run.style.font.size:
            return run.style.font.size.pt
    except Exception:
        pass
    return style_size(para.style, dflt)


def set_rfonts(rfonts):
    for k in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn("w:" + k), None)
    cur = (rfonts.get(qn("w:ascii")) or "").lower()
    if cur in KEEP_FONTS:
        return
    for k in ("ascii", "hAnsi", "eastAsia"):
        rfonts.set(qn("w:" + k), FONT)


def fix_fonts(doc):
    root_parts = [doc.styles.element, doc.element]
    for sec in doc.sections:
        for hf in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer):
            try:
                root_parts.append(hf._element)
            except Exception:
                pass
    for root in root_parts:
        for rpr in root.iter(qn("w:rPr")):
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
            set_rfonts(rf)
    # 既定（docDefaults）にも
    dd = doc.styles.element.find(qn("w:docDefaults"))
    if dd is not None:
        rpr = dd.find(qn("w:rPrDefault"))
        if rpr is not None and rpr.find(qn("w:rPr")) is not None:
            r = rpr.find(qn("w:rPr")); rf = r.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts"); r.insert(0, rf)
            set_rfonts(rf)


def iter_blocks(doc):
    """(種類, Paragraph or Table, 参考文献以降か)"""
    refs = False
    for el in doc.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            p = Paragraph(el, doc)
            if REF_HEAD.match(p.text.strip()):
                refs = True
            yield "p", p, refs
        elif tag == "tbl":
            yield "t", Table(el, doc), refs


def table_paras(tbl):
    seen = set()   # 要素そのものを持つ（id() は使い回されて取りこぼす）
    for tr in tbl._tbl.tr_lst:
        for tc in tr.tc_lst:
            if tc in seen:
                continue
            seen.add(tc)
            for p in _Cell(tc, tbl).paragraphs:
                yield p


def fix_sizes(doc):
    dflt = default_size(doc)
    body = collections.Counter()
    for kind, obj, refs in iter_blocks(doc):
        if kind == "p" and not refs and not obj.style.name.lower().startswith(("heading", "title", "見出し", "表題")):
            for r in obj.runs:
                if r.text.strip() and not (r.font.superscript or r.font.subscript):
                    body[round(eff_size(r, obj, dflt) * 2) / 2] += len(r.text.strip())
    mode = body.most_common(1)[0][0] if body else BODY_TARGET
    k = BODY_TARGET / mode if mode < BODY_MIN else 1.0
    changed = collections.Counter()

    def bump(r, p, floor, scale=1.0):
        if not r.text.strip() or r.font.superscript or r.font.subscript:
            return
        s = eff_size(r, p, dflt)
        ns = max(s * scale, floor) if s < 14 else s * (scale if scale > 1 else 1)  # 見出しも比例
        ns = round(ns * 2) / 2
        if abs(ns - s) >= 0.25:
            r.font.size = Pt(ns); changed["runs"] += 1

    for kind, obj, refs in iter_blocks(doc):
        if kind == "p":
            heading = obj.style.name.lower().startswith(("heading", "title", "見出し", "表題"))
            if refs:
                for r in obj.runs: bump(r, obj, REF_MIN)
            elif heading:
                if k > 1:
                    for r in obj.runs: bump(r, obj, 0, k)
            else:
                for r in obj.runs: bump(r, obj, BODY_MIN, k)
        else:
            ncol = len(obj._tbl.tblGrid.findall(qn("w:gridCol")))
            floor = 9.5 if ncol <= 4 else 9.0
            if refs:
                floor = REF_MIN
            for p in table_paras(obj):
                for r in p.runs: bump(r, p, floor)
    return mode, k, changed["runs"]


def fix_figures(doc):
    sec = doc.sections[0]
    maxw = (sec.page_width - sec.left_margin - sec.right_margin)
    maxh = (sec.page_height - sec.top_margin - sec.bottom_margin) * 0.75
    n = 0
    for ext in doc.element.body.iter(qn("wp:extent")):
        cx, cy = int(ext.get("cx")), int(ext.get("cy"))
        if cx / 360000 >= FIG_MIN_CM or cx < 360000 * 3:   # 3cm 未満はアイコン等
            continue
        target = min(int(FIG_CM * 360000), int(maxw))
        f = target / cx
        if cy * f > maxh:
            f = maxh / cy
        if f <= 1.02:
            continue
        ext.set("cx", str(int(cx * f))); ext.set("cy", str(int(cy * f)))
        inline = ext.getparent()
        for a in inline.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}ext"):
            if a.get("cx"):
                a.set("cx", str(int(int(a.get("cx")) * f))); a.set("cy", str(int(int(a.get("cy")) * f)))
        n += 1
    return n


def fix_table_rows(doc):
    for tbl in doc.tables:
        for i, tr in enumerate(tbl._tbl.tr_lst):
            trPr = tr.get_or_add_trPr()
            if trPr.find(qn("w:cantSplit")) is None:
                trPr.append(OxmlElement("w:cantSplit"))
            if i == 0 and len(tbl._tbl.tr_lst) > 2:
                if trPr.find(qn("w:tblHeader")) is None:
                    trPr.append(OxmlElement("w:tblHeader"))
                for tc in tr.tc_lst:
                    for p in _Cell(tc, tbl).paragraphs:
                        p.paragraph_format.keep_with_next = True


def word_finish(path, intended, pdf=True):
    import win32com.client, pythoncom
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False; word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(os.path.abspath(path))
        try:
            for toc in d.TablesOfContents:
                toc.Update()
            d.Fields.Update()
        except Exception:
            pass
        d.Save(); d.Close(False)
    finally:
        word.Quit()
    left = FT.word_verify(path, intended, pdf=pdf)
    return left


def process(path, pdf=True):
    folder, name = os.path.split(path)
    stem = os.path.splitext(name)[0]
    bak = os.path.join(folder, "_旧版"); os.makedirs(bak, exist_ok=True)
    shutil.copyfile(path, os.path.join(bak, stem + "_見やすさ調整前.docx"))
    if os.path.exists(os.path.join(folder, stem + ".pdf")):
        shutil.copyfile(os.path.join(folder, stem + ".pdf"), os.path.join(bak, stem + "_見やすさ調整前.pdf"))
    doc = Document(path)
    fix_fonts(doc)
    mode, k, nruns = fix_sizes(doc)
    nfig = fix_figures(doc)
    fix_table_rows(doc)
    doc.save(path)
    rep, intended = FT.fit_docx(path, path)
    left = word_finish(path, intended, pdf=pdf)
    return dict(body_mode=mode, scale=round(k, 2), runs=nruns, figs=nfig, tables=len(rep), wraps_left=len(left))


if __name__ == "__main__":
    pdf = "--no-pdf" not in sys.argv
    for p in [a for a in sys.argv[1:] if not a.startswith("--")]:
        try:
            print(os.path.basename(p), process(p, pdf), flush=True)
        except Exception as e:
            print(os.path.basename(p), "ERROR", repr(e), flush=True)


def reprocess(path, pdf=True):
    """控えを作らずにもう一度かける（調整のやり直し用）。"""
    doc = Document(path)
    fix_fonts(doc); mode, k, nruns = fix_sizes(doc); fix_figures(doc); fix_table_rows(doc)
    doc.save(path)
    rep, intended = FT.fit_docx(path, path)
    left = word_finish(path, intended, pdf=pdf)
    return dict(runs=nruns, wraps_left=len(left))
