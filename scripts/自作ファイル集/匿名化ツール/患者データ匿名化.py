# -*- coding: utf-8 -*-
"""
患者データ匿名化ツール(不可逆マスク方式)

Excelファイル(.xlsx / .xlsm)内の個人情報列をヘッダー名から自動検出し、
復元不可能なマスク文字列に置き換えた新しいファイルを出力します。
元のファイルは一切変更しません。

使い方:
    python 患者データ匿名化.py ファイル1.xlsx [ファイル2.xlsx ...]
    (または「ここにドラッグ.bat」にExcelファイルをドラッグ&ドロップ)

出力:
    元ファイルと同じフォルダに「元ファイル名_匿名化.xlsx」を作成
"""

import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

# ---------------------------------------------------------------
# 匿名化対象の列(ヘッダー名のパターン)とマスク後の文字列
# 必要に応じてパターンを追加してください
# ---------------------------------------------------------------
TARGET_COLUMNS = [
    # (分類名, ヘッダー名の正規表現, マスク文字列)
    ("氏名",     r"(患者氏?名|氏名|名前|患者名|漢字氏名|様氏名)", "匿名"),
    ("カナ氏名", r"(フリガナ|ふりがな|カナ氏名|氏名カナ|カナ|シメイ)", "トクメイ"),
    ("患者ID",   r"(患者ID|患者番号|カルテ番号|カルテNo|診察券番号|診察券No|受診者番号|会員番号|^ID$|^No\.?$)", "*****"),
    ("生年月日", r"(生年月日|誕生日|出生日|DOB)", "****/**/**"),
    ("住所",     r"(住所|所在地|居住地|郵便番号|〒)", "匿名"),
    ("電話番号", r"(電話|TEL|携帯|FAX|ファックス|連絡先)", "***-****-****"),
    ("保険情報", r"(保険者番号|被保険者|保険証|記号番号|公費負担者番号|受給者番号|マイナンバー|個人番号)", "*****"),
    ("メール",   r"(メール|Mail|E-?mail)", "*****"),
    ("勤務先",   r"(勤務先|職業|所属|学校名)", "匿名"),
    ("緊急連絡", r"(緊急連絡|キーパーソン|家族氏名|保護者|続柄)", "匿名"),
]

# 自由記載欄などに紛れた電話番号・メールアドレスも念のため全セル走査して消す
RE_PHONE = re.compile(r"0\d{1,4}-\d{1,4}-\d{3,4}|0\d{9,10}")
RE_EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")

HEADER_SEARCH_ROWS = 10   # 先頭から何行目までをヘッダー行の候補として探すか


def normalize(text):
    """全角/半角・大文字小文字の揺れを吸収する"""
    return unicodedata.normalize("NFKC", str(text)).strip().lower()


def find_header_row(ws):
    """対象パターンに最も多くヒットする行をヘッダー行とみなす"""
    best_row, best_hits = None, 0
    for row in ws.iter_rows(min_row=1, max_row=min(HEADER_SEARCH_ROWS, ws.max_row)):
        hits = 0
        for cell in row:
            if cell.value is None:
                continue
            header = normalize(cell.value)
            for _, pattern, _ in TARGET_COLUMNS:
                if re.search(normalize(pattern), header, re.IGNORECASE):
                    hits += 1
                    break
        if hits > best_hits:
            best_row, best_hits = row[0].row, hits
    return best_row


def anonymize_sheet(ws):
    """1シートを匿名化し、(検出列リスト, 置換セル数, 走査置換数) を返す"""
    header_row = find_header_row(ws)
    detected = []   # (列文字, ヘッダー名, 分類, マスク文字列)
    masked_cells = 0

    if header_row is not None:
        for cell in ws[header_row]:
            if cell.value is None:
                continue
            header = normalize(cell.value)
            for category, pattern, mask in TARGET_COLUMNS:
                if re.search(normalize(pattern), header, re.IGNORECASE):
                    detected.append((cell.column_letter, str(cell.value), category, mask))
                    break

        for col_letter, _, _, mask in detected:
            for row in range(header_row + 1, ws.max_row + 1):
                cell = ws[f"{col_letter}{row}"]
                if cell.value is not None and str(cell.value).strip() != "":
                    cell.value = mask
                    masked_cells += 1

    # 全セル走査: 自由記載欄の電話番号・メールアドレスを消す
    swept = 0
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                new_val = RE_PHONE.sub("***-****-****", cell.value)
                new_val = RE_EMAIL.sub("*****", new_val)
                if new_val != cell.value:
                    cell.value = new_val
                    swept += 1

    return detected, masked_cells, swept


def anonymize_file(path):
    path = Path(path)
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        print(f"[スキップ] {path.name}: xlsx / xlsm 以外は対象外です")
        return

    keep_vba = path.suffix.lower() == ".xlsm"
    wb = openpyxl.load_workbook(path, keep_vba=keep_vba)

    print(f"\n=== {path.name} ===")
    total_masked = 0
    for ws in wb.worksheets:
        detected, masked, swept = anonymize_sheet(ws)
        total_masked += masked + swept
        if detected:
            print(f"  シート「{ws.title}」: {masked}セルをマスク")
            for col, header, category, _ in detected:
                print(f"    {col}列「{header}」({category})")
        else:
            print(f"  シート「{ws.title}」: 対象列は検出されませんでした")
        if swept:
            print(f"    ほか全セル走査で電話番号/メールを {swept}セル置換")

    # 出力ファイル名(既存なら連番を付けて上書き回避)
    out = path.with_name(f"{path.stem}_匿名化{path.suffix}")
    n = 2
    while out.exists():
        out = path.with_name(f"{path.stem}_匿名化({n}){path.suffix}")
        n += 1
    wb.save(out)
    print(f"  → 出力: {out.name}(計{total_masked}セル置換)")

    if total_masked == 0:
        print("  ※ 1件も置換されていません。ヘッダー名が想定パターンと違う可能性があります。")


def main():
    if len(sys.argv) < 2:
        print("使い方: python 患者データ匿名化.py <Excelファイル> [複数指定可]")
        print("または「ここにドラッグ.bat」にファイルをドラッグ&ドロップしてください。")
        return
    for arg in sys.argv[1:]:
        try:
            anonymize_file(arg)
        except Exception as e:
            print(f"[エラー] {arg}: {e}")
    print("\n完了しました。元ファイルは変更されていません。")


if __name__ == "__main__":
    main()
