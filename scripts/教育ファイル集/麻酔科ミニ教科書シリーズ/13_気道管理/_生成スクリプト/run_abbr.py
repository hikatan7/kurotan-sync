# -*- coding: utf-8 -*-
"""1 つの docx に略語処理を適用する。

  python run_abbr.py <入力docx> <巻番号2桁> [出力docx]
"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import abbr_dict
from apply_abbr import (ptext, pstyle, is_body_para, expand_in_para,
                        make_para, page_break_para, CHAP_RE, TOC_STYLES)

NAVY = "0A1830"
HEADFILL = "2F4A6B"


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def process(path, vol, out=None):
    out = out or path
    doc = Document(path)
    body = doc.element.body
    D = abbr_dict.dict_for(vol)
    # 長い略語から処理して NT-proBNP が BNP に壊されないようにする
    keys = sorted(D.keys(), key=lambda k: (-len(k), k))

    tops = [e for e in body if e.tag in (qn("w:p"), qn("w:tbl"))]
    # 書式の見本になる本文段落
    sample = next((e for e in tops if is_body_para(e)), None)
    if sample is None:
        return None

    used = set()          # 巻全体で実際に本文に出た略語
    done_vol = set()      # COMMON_ONCE 用
    done_chap = set()     # 章内で展開済み
    n_exp = 0

    for e in tops:
        if e.tag != qn("w:p"):
            continue
        t = ptext(e).strip()
        if pstyle(e) not in TOC_STYLES and CHAP_RE.match(t):
            done_chap = set()
        if not is_body_para(e):
            # 使用実績の記録だけは全文から拾う
            for a in keys:
                if re.search(r"(?<![A-Za-z0-9\-])" + re.escape(a)
                             + r"(?![A-Za-z0-9\-])", t):
                    used.add(a)
            continue
        for a in keys:
            if not re.search(r"(?<![A-Za-z0-9\-])" + re.escape(a)
                             + r"(?![A-Za-z0-9\-])", ptext(e)):
                continue
            used.add(a)
            once = a in abbr_dict.COMMON_ONCE
            if once and a in done_vol:
                continue
            if not once and a in done_chap:
                continue
            eng, jp = D[a]
            tx = ptext(e)
            # すでに説明がある段落は、そこを初出とみなして触らない
            if re.search(re.escape(a) + r"[（(：:]", tx) or eng.lower() in tx.lower():
                done_vol.add(a)
                done_chap.add(a)
                continue
            if expand_in_para(e, a, eng, jp):
                n_exp += 1
                done_vol.add(a)
                done_chap.add(a)

    # 表の中だけに出る略語も一覧には載せる
    for tb in body.iter(qn("w:tbl")):
        t = ptext(tb)
        for a in keys:
            if re.search(r"(?<![A-Za-z0-9\-])" + re.escape(a)
                         + r"(?![A-Za-z0-9\-])", t):
                used.add(a)

    if not used:
        doc.save(out)
        return (0, 0)

    # ---- 巻頭の略語一覧を作る ----
    anchor = None
    for e in tops:
        if e.tag != qn("w:p") or pstyle(e) in TOC_STYLES:
            continue
        t = ptext(e).strip()
        if re.match(r"^(はじめに|第\s*[IVX０-９0-9]+\s*[部章])", t):
            anchor = e
            break
    if anchor is None:
        anchor = sample

    rows = sorted(used, key=lambda a: a.upper())
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    for i, h in enumerate(("略語", "英語", "日本語")):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade_cell(hdr[i], HEADFILL)
    for a in rows:
        eng, jp = D[a]
        cs = tbl.add_row().cells
        for i, v in enumerate((a, eng, jp)):
            cs[i].text = ""
            p = cs[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.12
            r = p.add_run(v)
            r.font.size = Pt(8.5)
            if i == 0:
                r.bold = True
    for row in tbl.rows:
        row.cells[0].width = Cm(2.6)
        row.cells[1].width = Cm(7.6)
        row.cells[2].width = Cm(6.6)

    els = [page_break_para(sample),
           make_para(sample, "本書で使う略語", size=14, bold=True, color=NAVY),
           make_para(sample,
                     "本文では、章ごとの初出で「略語（英語表記：日本語表記）」の形に展開している。"
                     "以下は本書に出てくる略語の一覧である。", size=10.5),
           tbl._tbl,
           page_break_para(sample)]
    prev = None
    for x in els:
        if prev is None:
            anchor.addprevious(x)
        else:
            prev.addnext(x)
        prev = x

    doc.save(out)
    return (n_exp, len(rows))


if __name__ == "__main__":
    src = sys.argv[1]
    vol = sys.argv[2]
    dst = sys.argv[3] if len(sys.argv) > 3 else None
    r = process(src, vol, dst)
    print(os.path.basename(src), "→ 展開", r[0], "件／一覧", r[1], "語")
