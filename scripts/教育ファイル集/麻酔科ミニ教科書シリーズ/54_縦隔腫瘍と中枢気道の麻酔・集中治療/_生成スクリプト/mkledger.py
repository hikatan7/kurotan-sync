# -*- coding: utf-8 -*-
"""確定した PMID の書誌を PubMed から取り、台帳 xlsx と refs54.py を作る。"""
import json, os, sys, time, urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
PMIDS = open(os.path.join(HERE, "pmids.txt"), encoding="utf-8").read().split()

E = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi?db=pubmed&retmode=json&id="
recs = {}
for i in range(0, len(PMIDS), 50):
    ch = PMIDS[i:i + 50]
    with urllib.request.urlopen(E + ",".join(ch), timeout=90) as r:
        d = json.load(r)["result"]
    for u in ch:
        recs[u] = d.get(u, {})
    time.sleep(0.4)

AUTHOR_OVERRIDE = {
    # esummary が著者を返さない団体著作。手で補う。
    "26580836": "American Society of Anesthesiologists Task Force on Obstetric Anesthesia",
    "29334501": "American Society of Anesthesiologists Task Force on Moderate Procedural Sedation and Analgesia",
    "36629465": "American Society of Anesthesiologists Task Force on Preoperative Fasting",
}

EXTRA = {
    "APSF2025": "Raghavan KC, Paquin JR. Anesthetic considerations for the management of patients with anterior mediastinal mass. APSF Newsletter. 2025 Sep 2. https://www.apsf.org/article/anesthetic-considerations-for-the-management-of-patients-with-anterior-mediastinal-mass/",
    "CHOUDHURY2026": "Choudhury M, et al. Management of post-operative anterior mediastinal mass patients: a review. J Card Crit Care TSS. 2026;10(3):173-181. doi:10.25259/JCCC_30_2025",
    "AABIP2023": "American Association for Bronchology and Interventional Pulmonology. Statement on fundamental principles and safety of bronchoscopic electrocautery and argon plasma coagulation. 2023 Dec 1. https://www.aabronchology.org/assets/docs/AdvocacyStatement/FINAL_Advocacy%20statement_Dec%202023.pdf",
    "FDA2005": "U.S. Food and Drug Administration. FDA Public Health Notification: complications from metallic tracheal stents in patients with benign airway disorders. 2005 Jul 29.",
}

rows = []
for u in PMIDS:
    r = recs.get(u) or {}
    aus = [a["name"] for a in r.get("authors", []) if a.get("authtype") != "CollectiveName"]
    if not aus and r.get("authors"):
        aus = [r["authors"][0]["name"]]
    au = aus[0] + (", et al." if len(aus) > 1 else "") if aus else ""
    au = AUTHOR_OVERRIDE.get(u, au)
    doi = ""
    for aid in r.get("articleids", []):
        if aid.get("idtype") == "doi":
            doi = aid.get("value", "")
    rows.append({
        "pmid": u, "author": au, "authors_all": "; ".join(aus[:6]),
        "title": r.get("title", "").rstrip("."),
        "journal": r.get("source", ""), "year": str(r.get("pubdate", ""))[:4],
        "volume": r.get("volume", ""), "issue": r.get("issue", ""),
        "pages": r.get("pages", ""), "doi": doi,
    })

miss = [r["pmid"] for r in rows if not r["title"]]
print("MISSING:", miss)

with open(os.path.join(HERE, "refs54.py"), "w", encoding="utf-8") as f:
    f.write("# -*- coding: utf-8 -*-\n")
    f.write('"""文献台帳（PubMed esummary から生成）。手で編集しない。"""\n')
    f.write("REFS = {\n")
    for r in rows:
        cite = "%s. %s. %s. %s" % (r["author"].rstrip("."), r["title"].rstrip("."), r["journal"].rstrip("."), r["year"])
        if r["volume"]:
            cite += ";%s" % r["volume"]
            if r["issue"]:
                cite += "(%s)" % r["issue"]
            if r["pages"]:
                cite += ":%s" % r["pages"]
        cite += ". PMID: %s" % r["pmid"]
        if r["doi"]:
            cite += ". doi:%s" % r["doi"]
        f.write("    %r: %r,\n" % (r["pmid"], cite))
    for k, v in EXTRA.items():
        f.write("    %r: %r,\n" % (k, v))
    f.write("}\n")

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "文献台帳"
    head = ["No", "PMID", "第一著者", "共著者", "タイトル", "誌名", "年",
            "巻", "号", "頁", "DOI"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = openpyxl.styles.PatternFill("solid", fgColor="2F4A6B")
    for i, r in enumerate(rows, 1):
        ws.append([i, r["pmid"], r["author"], r["authors_all"], r["title"],
                   r["journal"], r["year"], r["volume"], r["issue"],
                   r["pages"], r["doi"]])
    for col, w in zip("ABCDEFGHIJK", (5, 11, 20, 34, 72, 22, 7, 7, 6, 12, 30)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    wb.save(os.path.join(HERE, "..", "縦隔腫瘍と中枢気道の麻酔・集中治療_SourceLedger.xlsx"))
    print("ledger saved")
except Exception as e:
    print("xlsx skipped:", e)
sys.stdout.reconfigure(encoding="utf-8")
print("n =", len(rows))
