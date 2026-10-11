# -*- coding: utf-8 -*-
"""⑱ 専攻医版30頁（コラム入り）を組む。"""
import os, re, sys
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm
import bookkit
from bookkit import (drop_orphan_blanks, new_document, set_footer, para, heading,
                     figure, table, bullets, summary, note, lead, column, references,
                     shaded, left_bar, _run, NAVY, GREY, TEAL, BLUE, INK, LINE_TIGHT)
import volno, refs18s, c18s_p1, c18s_p2

bookkit.REFS = dict(refs18s.REFS)
FIGDIR = r"C:/Users/junch/OneDrive/_旧デスクトップ/教育ファイル集/麻酔科ミニ教科書シリーズ/18_呼吸管理と分離肺換気/図版"
OUT = os.path.join(HERE, "呼吸管理と分離肺換気_専攻医版30頁.docx")
DROP = set(os.environ.get("DROPFIG", "").split(","))
BLOCKS = [b for b in volno.fix(c18s_p1.BLOCKS + c18s_p2.BLOCKS)
          if not (b[0] == "fig" and b[1] in DROP)]
FIGW = float(os.environ.get("FIGW", "13.5"))


def case(doc, title, pres, answers):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    shaded(p, "EAF1F8"); left_bar(p, "174682")
    _run(p, "◆ " + title, size=11, bold=True, color=BLUE)
    q = doc.add_paragraph()
    q.paragraph_format.space_after = Pt(3)
    q.paragraph_format.line_spacing = LINE_TIGHT
    q.paragraph_format.keep_with_next = True
    shaded(q, "EAF1F8"); left_bar(q, "174682")
    _run(q, pres, size=10, color=INK)
    r = doc.add_paragraph()
    r.paragraph_format.space_after = Pt(1)
    r.paragraph_format.keep_with_next = True
    _run(r, "▷ 答えと考え方", size=10, bold=True, color=BLUE)
    for i, a in enumerate(answers):
        s = doc.add_paragraph()
        s.paragraph_format.space_after = Pt(2)
        s.paragraph_format.line_spacing = LINE_TIGHT
        s.paragraph_format.left_indent = Cm(0.55)
        s.paragraph_format.first_line_indent = Cm(-0.45)
        s.paragraph_format.keep_with_next = (i < len(answers) - 1)
        _run(s, "・" + a, size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def build():
    doc = new_document()
    set_footer(doc, "麻酔科ミニ教科書シリーズ ⑱　別冊　呼吸管理と分離肺換気　専攻医版　　")
    para(doc, "麻酔科ミニ教科書シリーズ ⑱　別冊", size=11, bold=True, color=TEAL,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc, "呼吸管理と分離肺換気　── 専攻医版 ──", size=19, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    para(doc, "肺を分ける理由・道具・確認・設定 ── 設定値から始めない", size=12,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "2026 年 9 月 24 日　静岡市立静岡病院　麻酔科　玉里順子", size=10, color=GREY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    n = {"fig": 0, "tbl": 0, "col": 0, "case": 0}
    for b in BLOCKS:
        k = b[0]
        if k == "h2": heading(doc, b[1], 2)
        elif k == "h3": heading(doc, b[1], 3)
        elif k == "p": para(doc, b[1])
        elif k == "lead": lead(doc, b[1])
        elif k == "bul": bullets(doc, b[1])
        elif k == "fig":
            n["fig"] += 1
            cap = re.sub(r"^図\d+", "図%d" % n["fig"], b[2])
            figure(doc, os.path.join(FIGDIR, b[1]), cap, width_cm=FIGW)
        elif k == "tbl": table(doc, None, b[1], b[2], size=8.5); n["tbl"] += 1
        elif k == "sum": summary(doc, b[1], b[2])
        elif k == "note": note(doc, b[1], b[2])
        elif k == "col":
            n["col"] += 1; column(doc, n["col"], b[1], b[2])
        elif k == "case": case(doc, b[1], b[2], b[3]); n["case"] += 1
        elif k == "refs": n["refs"] = references(doc, size=8.0)
        else: raise ValueError(k)
    drop_orphan_blanks(doc)
    doc.save(OUT)
    return n

if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    print(build())
