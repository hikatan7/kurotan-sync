# -*- coding: utf-8 -*-
"""既刊の docx（控え）から本文を取り出し、bookkit のブロック列に戻す。

PDF から戻すより確実。docx には次のものがそのまま残っている。

    ・見出しの階層  … 段落の「フォントサイズ＋太字」で決まる（見出しスタイルは未使用）
    ・枠            … 1列1行の表。1つ目のセルが 【タグ】 見出し\\n本文… の形
    ・表            … 2列以上の表。1行目が見出し行
    ・図            … 段落内のインライン画像。直後の小さい文字がキャプション

本文は書き換えない。判断するのは「どの段落が何なのか」の割り当てだけ。
"""
import os
import re
import collections

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

BOX = re.compile(r"^\s*【([^】]+)】\s*(.*)$")
# 部の番号は「I」だけでなくローマ数字1文字（Ⅰ Ⅱ …）で書かれている巻がある
PART = re.compile(r"^第\s*[IVXivxⅠ-Ⅻ]+\s*部")
CHAP = re.compile(r"^第\s*\d+\s*章")
APPX = re.compile(r"^付\s*録\s*[A-Z0-9]")
FIGCAP = re.compile(r"^図\s*[\d-]+")
TBLCAP = re.compile(r"^表\s*[\d A-Z-]+")
SUBH = re.compile(r"^[■◆]\s*")
BULLET = re.compile(r"^[・‣▸]\s*")
HSTYLE = re.compile(r"^(?:Heading|見出し)\s*(\d)")
# 「Chapter 1」だけの行を章番号のラベルとして使っている巻（⑦）がある
CHAPLBL = re.compile(r"^(?:Chapter|CHAPTER|Chap)\s*(\d+)\s*$")


def _size(p, base):
    """段落の代表フォントサイズ。指定が無ければ既定値。"""
    for r in p.runs:
        if r.text.strip() and r.font.size:
            return round(r.font.size.pt, 1)
    if p.style and p.style.font.size:
        return round(p.style.font.size.pt, 1)
    return base


def _bold(p):
    """先頭の実体のある run が太字か。"""
    for r in p.runs:
        if r.text.strip():
            return bool(r.bold)
    return False


# 枠の種類。巻によって 【Take Home】 とも「Take Home Message」とも書いてある
TAGS = ["Take Home Message", "Take Home", "Clinical Pearl", "Pearl", "Pitfall",
        "Evidence", "Mechanism", "History", "Column", "コラム", "Case",
        "Q & A", "Q&A", "計算してみる", "やってみる"]
KIND = {"Take Home": "th", "Take Home Message": "th",
        "Pitfall": "pf", "Evidence": "ev"}
_TAGRE = re.compile(r"^\s*(?:【\s*)?(" + "|".join(
    re.escape(t) for t in TAGS) + r")(?:\s*】)?\s*(?:[—―–\-:：]\s*)?(.*)$")


def _box_tag(line):
    """枠の1行目から (タグ, 題) を取り出す。タグが無ければ (None, None)。"""
    m = _TAGRE.match(line)
    if not m:
        m2 = BOX.match(line)
        if m2:
            return m2.group(1), m2.group(2).strip()
        return None, None
    return m.group(1), m.group(2).strip()


# タグは無いが、題が「第15章のまとめ」なら章のまとめ
SUMTITLE = re.compile(r"(?:章|部|節|本章|ここ)?の?(?:まとめ|要点)$|^まとめ$|^要点$")


def _is_title(s):
    """枠の題らしい行か。箇条書きや長い文は題ではない。"""
    return (len(s) <= 46 and not BULLET.match(s)
            and not re.match(r"^[①-⑳\d]", s) and not s.endswith("。"))


def _is_toc(p):
    """古い目次フィールドの行。見出しと紛らわしいので数えない。"""
    n = (p.style.name or "").lower()
    return n.startswith("toc") or n.startswith("目次")


# 「● Pitfall　専攻医が踏みやすい落とし穴。」のような枠の凡例。
# 枠そのものを廃止したので、説明だけ残ると宙に浮く。
LEGEND = re.compile(r"^[●○・■◆\-]?\s*(?:" + "|".join(
    re.escape(t) for t in ["Take Home Message", "Take Home", "Clinical Pearl",
                           "Pitfall", "Evidence Update", "Evidence", "Mechanism",
                           "History", "Column / History", "Column", "Case",
                           "Q & A", "計算してみる"]) + r")\s*[　 ]")


def _is_legend(t):
    return len(t) <= 90 and bool(LEGEND.match(t))


def _headish(t):
    """見出しになりうる短い行か（本文の「第2章で全体像を…」を弾く）。"""
    return len(t) <= 60 and "。" not in t


def _cell_lines(cell):
    out = []
    for p in cell.paragraphs:
        t = p.text.strip()
        if t:
            out.append(t)
    return out


def _rows(t):
    """表を行列に直す。結合セルは同じ文字列が繰り返されるので畳む。"""
    out = []
    for r in t.rows:
        row, prev = [], None
        for c in r.cells:
            v = c.text.strip().replace("\n", " ")
            row.append("" if v == prev and v else v)
            prev = v
        out.append(row)
    return [r for r in out if any(r)]


def _images(par):
    """段落に入っている画像の rId を、出てくる順に返す。"""
    ids = []
    for b in par._p.iter():
        if b.tag == qn("a:blip"):
            rid = b.get(qn("r:embed"))
            if rid:
                ids.append(rid)
    return ids


def profile(doc):
    """本文サイズと、部・章の見出しサイズを実測する。

    見出しの大きさは巻ごとに違う（①は 部20・章16・節13）。
    決め打ちすると節が章に混ざるので、「第X部」「第N章」と書いてある段落の
    実際のサイズを測り、それを基準にする。
    """
    base = 10.5
    if doc.styles["Normal"].font.size:
        base = round(doc.styles["Normal"].font.size.pt, 1)
    cnt = collections.Counter()
    h1s, h2s = collections.Counter(), collections.Counter()
    styles = collections.Counter()
    part_in_h1 = False
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t or _is_toc(p):
            continue
        m = HSTYLE.match(p.style.name or "")
        if m:
            styles[int(m.group(1))] += 1
            if m.group(1) == "1" and PART.match(t):
                part_in_h1 = True
        sz, bd = _size(p, base), _bold(p)
        cnt[(sz, bd)] += 1
        if not bd or sz <= base or not _headish(t):
            continue
        if PART.match(t):
            h1s[sz] += 1
        elif CHAP.match(t) or APPX.match(t):
            h2s[sz] += 1
    h2 = max(h2s, key=lambda k: (h2s[k], k)) if h2s else base + 4.0
    h1 = max(h1s, key=lambda k: (h1s[k], k)) if h1s else h2 + 4.0
    if h1 <= h2:
        h1 = h2 + 4.0

    # ④のように本物の見出しスタイルを使っている巻は、そちらを優先する。
    # 部が無い巻では「見出し1 ＝ 章」なので、1段ずらして受ける。
    lv = {}
    if sum(styles.values()) >= 10:
        off = 0 if part_in_h1 else 1
        for k in styles:
            lv[k] = min(3, k + off)
    # ⑪のように章と節が同じ大きさの巻がある。その場合は「第N章／付録X」と
    # 書いてあるものだけを章として扱わないと、節まで章になってしまう。
    n_at_h2 = sum(n for (sz, bd), n in cnt.items() if bd and abs(sz - h2) < 0.3)
    need_pat = n_at_h2 > max(h2s.values(), default=0) * 1.5 + 4

    return {"body": base, "h1": h1, "h2": h2, "counts": cnt, "levels": lv,
            "has_part": bool(h1s) or part_in_h1, "h2_pattern": need_pat}


def convert(path, figdir):
    """(blocks, prof) を返す。図は figdir に f01.png… として書き出す。"""
    doc = docx.Document(path)
    prof = profile(doc)
    base = prof["body"]
    if figdir:
        os.makedirs(figdir, exist_ok=True)

    blocks = []
    buf, buf_kind = [], None
    nfig = 0
    label = ""

    def flush():
        nonlocal buf, buf_kind
        if not buf:
            return
        if buf_kind == "ul":
            blocks.append(("ul", [BULLET.sub("", x).strip() for x in buf]))
        else:
            blocks.append(("p", "".join(buf)))
        buf, buf_kind = [], None

    def push_box(cell):
        lines = _cell_lines(cell)
        if not lines:
            return
        tag, title = _box_tag(lines[0])
        body = lines[1:]
        if tag is None:
            clean = [BULLET.sub("", x).strip() for x in body]
            if SUMTITLE.search(lines[0]) and clean:
                blocks.append(("th", lines[0], clean))
                return
            # タグの無い囲み（④の「原則」「本書の中心命題」など）は注意書きに寄せる
            blocks.append(("note", lines[0], clean or [lines[0]]))
            return
        # ⑦のように、タグ行の次の行が枠の題になっている巻がある
        if not title and body and _is_title(body[0]):
            title, body = body[0], body[1:]
        body = [BULLET.sub("", x).strip() for x in body]
        kind = KIND.get(tag)
        if kind:
            blocks.append((kind, title or tag, body))
        else:
            blocks.append(("colx", tag, title, body))

    for ch in doc.element.body.iterchildren():
        tag = ch.tag.split("}")[-1]

        if tag == "tbl":
            flush()
            t = Table(ch, doc)
            if len(t.columns) == 1:
                push_box(t.rows[0].cells[0])
                continue
            rows = _rows(t)
            if len(rows) >= 2:
                blocks.append(("tbl", rows[0], rows[1:], None))
            continue

        if tag != "p":
            continue

        p = Paragraph(ch, doc)
        rids = _images(p)
        if rids:
            flush()
            for rid in rids:
                nfig += 1
                part = doc.part.related_parts[rid]
                ext = os.path.splitext(part.partname)[1] or ".png"
                fn = "f%02d%s" % (nfig, ext)
                if figdir:
                    with open(os.path.join(figdir, fn), "wb") as f:
                        f.write(part.blob)
                blocks.append(("fig", fn, ""))
            continue

        t = p.text.strip()
        if not t or _is_toc(p) or _is_legend(t):
            continue
        sz, bd = _size(p, base), _bold(p)

        # 「Chapter 1」だけの行は章番号のラベル。次の見出しの頭に付ける
        m = CHAPLBL.match(t)
        if m:
            flush()
            label = "第%s章　" % m.group(1)
            continue

        lvl = None
        m = HSTYLE.match(p.style.name or "")
        if m and prof["levels"]:
            lvl = prof["levels"].get(int(m.group(1)))
        elif bd and sz >= prof["h1"] - 0.5:
            lvl = 1
        elif bd and sz >= prof["h2"] - 0.5:
            lvl = 2
            if prof["h2_pattern"] and not (CHAP.match(t) or APPX.match(t)):
                lvl = 3
        elif bd and (sz >= base + 1.5 or SUBH.match(t)):
            lvl = 3
        # 大きさより文字のほうが確かなので、最後に上書きする。
        # ⑩は「第II部」だけ他の部より小さく、付録が部と同じ大きさだった。
        if lvl:
            if PART.match(t):
                lvl = 1
            elif CHAP.match(t) or APPX.match(t):
                lvl = 2
        # 部の無い巻（⑦など）では、大きい見出しも章として扱う。
        if lvl == 1 and not prof["has_part"]:
            lvl = 2
        if lvl:
            flush()
            t = re.sub(r"\s*\n\s*", "　", SUBH.sub("", t))
            if label and lvl <= 2:
                t, label = label + t, ""
                lvl = 2          # ラベルが付くのは章
            blocks.append(("h%d" % lvl, t))
            continue
        label = ""

        if sz <= base - 1.0:
            # 図表のキャプションと、表の注記・文献行
            if FIGCAP.match(t):
                for i in range(len(blocks) - 1, max(-1, len(blocks) - 3), -1):
                    if blocks[i][0] == "fig" and not blocks[i][2]:
                        blocks[i] = ("fig", blocks[i][1], t)
                        break
                continue
            if TBLCAP.match(t):
                continue
            if buf_kind == "p" and buf and not re.search(r"[。）\]]$", buf[-1]):
                buf[-1] += t
            else:
                flush()
                buf_kind, buf = "p", [t]
            continue

        # docx の段落は1つで完結しているので、PDF のように行を継ぎ足さない。
        # 続けて出てくる箇条書きだけ、1つの ul にまとめる。
        if BULLET.match(t):
            if buf_kind != "ul":
                flush()
                buf_kind = "ul"
            buf.append(t)
        else:
            flush()
            blocks.append(("p", t))

    flush()
    prof["abbrev"] = take_abbrev(blocks)
    return strip_front(blocks), prof


def take_abbrev(blocks):
    """巻頭の「本書で使う略語」の表を拾う。

    strip_front は第1章より前を全部落とすので、そのままだと
    略語一覧まで消えてしまう。本文に戻せるよう、ここで取っておく。
    """
    for i, b in enumerate(blocks):
        if b[0] != "tbl":
            continue
        head = "".join(b[1])
        if "略語" not in head and "Abbrev" not in head:
            continue
        if len(b[1]) >= 2 and len(b[2]) >= 10:
            return b
    return None


def strip_front(blocks):
    """表紙・略語表・はじめに・目次を落とし、本文の先頭から返す。

    目次の「第1章 …」は本文サイズで太字でないので見出しにならない。
    ここで探すのは、直後にまとまった段落が続く最初の「第1章」の見出し。
    """
    for i, b in enumerate(blocks):
        if b[0] != "h2" or not CHAP.match(b[1]):
            continue
        if not re.match(r"^第\s*1\s*章", b[1]):
            continue
        nxt = blocks[i + 1:i + 6]
        if any(x[0] == "p" and len(x[1]) > 80 for x in nxt):
            # 直前に「第 I 部」があれば、そこから返す（間に前書きが入る巻がある）
            for j in range(i - 1, max(-1, i - 8), -1):
                if blocks[j][0] == "h1":
                    return blocks[j:]
            return blocks[i:]
    return blocks


if __name__ == "__main__":
    import sys
    sys.stdout.reconfigure(encoding="utf-8")
    bl, pr = convert(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
    print("本文 %.1fpt / 部 %.1f / 章 %.1f" % (pr["body"], pr["h1"], pr["h2"]))
    print(collections.Counter(b[0] for b in bl))
    for b in bl[:40]:
        print("  %-5s %s" % (b[0], str(b[1])[:76]))
