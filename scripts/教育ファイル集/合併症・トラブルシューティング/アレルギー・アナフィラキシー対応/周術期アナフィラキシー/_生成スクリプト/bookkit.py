# -*- coding: utf-8 -*-
"""麻酔科ミニ教科書シリーズ 用 docx ビルダー

既刊（第11巻）の docx をテンプレートとして開き、本文を差し替える。
スタイル・ページ設定・フォントは既刊と同一になる。
"""
import copy
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

TEMPLATE = (r"C:\Users\junch\OneDrive\デスクトップ\教育ファイル集"
            r"\麻酔科ミニ教科書シリーズ\11_術前評価と麻酔計画\術前評価と麻酔計画.docx")

NAVY = RGBColor(0x0A, 0x18, 0x30)
GREY = RGBColor(0x55, 0x60, 0x6E)
TEAL = RGBColor(0x00, 0x7A, 0x8A)
GREEN = RGBColor(0x00, 0x5F, 0x4F)
BLUE = RGBColor(0x17, 0x46, 0x82)
PURPLE = RGBColor(0x4E, 0x38, 0x7A)
RED = RGBColor(0x9C, 0x2B, 0x26)
INK = RGBColor(0x22, 0x30, 0x3F)


def new_document():
    doc = Document(TEMPLATE)
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)
    return doc


def set_footer(doc, text):
    """テンプレート由来のフッター文言を差し替える（PAGE フィールドは保持）。"""
    for sec in doc.sections:
        f = sec.footer
        if f is None:
            continue
        for p in f.paragraphs:
            runs = list(p.runs)
            # PAGE フィールドが始まる位置より前のテキストランだけを書き換える
            cut = len(runs)
            for i, r in enumerate(runs):
                if r._r.findall(qn("w:fldChar")) or r._r.findall(qn("w:instrText")):
                    cut = i
                    break
            head = runs[:cut]
            if not any(r.text.strip() for r in head):
                continue
            for i, r in enumerate(head):
                r.text = text if i == 0 else ""
            break


# ---------------- 低レベル ----------------
def _run(p, text, size=None, bold=False, color=None, italic=False):
    r = p.add_run(text)
    if size:
        r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return r


def para(doc, text="", size=None, bold=False, color=None, align=None,
         space_before=None, space_after=None, line=1.37, indent=None,
         keep_with_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if space_after is not None:
        pf.space_after = Pt(space_after)
    if line:
        pf.line_spacing = line
    if indent is not None:
        pf.left_indent = Cm(indent)
    pf.keep_with_next = keep_with_next
    if text:
        _run(p, text, size=size, bold=bold, color=color)
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = NAVY
        if level == 1:
            r.font.size = Pt(17)
        elif level == 2:
            r.font.size = Pt(14)
        else:
            r.font.size = Pt(12)
    h.paragraph_format.keep_with_next = True
    return h


def toc_field(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    r._r.append(fld)
    r2 = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-3" \h \z \u '
    r2._r.append(instr)
    r3 = p.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(sep)
    r4 = p.add_run("（Word で開き Ctrl+A → F9 でページ番号が入ります）")
    r4.font.size = Pt(9)
    r4.font.color.rgb = GREY
    r5 = p.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r5._r.append(end)


def shaded(p, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    p._p.get_or_add_pPr().append(shd)


def left_bar(p, color, size=18):
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    pbdr.append(left)
    p._p.get_or_add_pPr().append(pbdr)


# ---------------- ブロック ----------------
def figure(doc, path, caption, width_cm=16.0):
    w, h = Image.open(path).size
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)
    _run(c, caption, size=9.5, bold=True, color=NAVY)


def boxed(doc, tag, title, lines, color, fill, bar):
    """【Take Home】【Pitfall】【Evidence】など"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.32
    p.paragraph_format.keep_with_next = True
    shaded(p, fill)
    left_bar(p, bar)
    _run(p, "【%s】" % tag, size=9.5, bold=True, color=color)
    _run(p, title, bold=True, color=color)
    for ln in lines:
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(2)
        q.paragraph_format.line_spacing = 1.32
        q.paragraph_format.left_indent = Cm(0.4)
        shaded(q, fill)
        left_bar(q, bar)
        _run(q, ln, size=10.5, color=INK)
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(10)


def column(doc, no, title, paras):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    shaded(p, "F3F0F8")
    left_bar(p, "4E387A")
    _run(p, "コラム%d　%s" % (no, title), bold=True, color=PURPLE)
    for t in paras:
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(3)
        q.paragraph_format.line_spacing = 1.35
        q.paragraph_format.left_indent = Cm(0.4)
        shaded(q, "F3F0F8")
        left_bar(q, "4E387A")
        _run(q, t, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(10)


def table(doc, caption, header, rows, widths=None, size=9.0):
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_before = Pt(10)
        c.paragraph_format.space_after = Pt(3)
        c.paragraph_format.keep_with_next = True
        _run(c, caption, size=9.5, bold=True, color=NAVY)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.autofit = True
    for i, htxt in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        pp = cell.paragraphs[0]
        pp.paragraph_format.space_after = Pt(1)
        pp.paragraph_format.line_spacing = 1.15
        _run(pp, htxt, size=size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shaded(pp, "2F4A6B")
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "2F4A6B")
        tcPr.append(shd)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            pp = cells[i].paragraphs[0]
            pp.paragraph_format.space_after = Pt(1)
            pp.paragraph_format.line_spacing = 1.15
            _run(pp, str(v), size=size, color=INK)
    if widths:
        for row in t.rows:
            for i, wcm in enumerate(widths):
                row.cells[i].width = Cm(wcm)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    return t


def bullets(doc, items, size=10.5, indent=0.5):
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.left_indent = Cm(indent)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        _run(p, "・" + it, size=size, color=INK)


def lead(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.35
    shaded(p, "F2F4F6")
    left_bar(p, "8D97A3", size=12)
    _run(p, text, size=10.5, color=RGBColor(0x3D, 0x4A, 0x5A))
