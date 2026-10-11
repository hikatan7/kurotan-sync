# -*- coding: utf-8 -*-
"""付録として本の途中に参考文献がある本で使う。

applycites と同じことをするが、生成した番号つき一覧を末尾に足すのではなく、
もとの参考文献の付録があった場所に差し込む。付録の記号（A・B・C…）が
飛ばないようにするため。

    python applycites_inplace.py <docx> <台帳> <シート> <cites json> <付録見出し> [注記]
"""
import io, json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import citedocx, applycites
from docx import Document

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def find_section(doc, head, end_head=None):
    """見出しの段落番号と、その節が終わる段落番号を返す。

    目次にも同じ文字列があるので、最初の一致ではなく本文側（最後の一致）を採る。
    これを間違えて本文を丸ごと消した事故があった。
    """
    ps = doc.paragraphs
    hits = [i for i, p in enumerate(ps) if p.text.strip() == head]
    if not hits:
        hits = [i for i, p in enumerate(ps) if p.text.strip().startswith(head)]
    if not hits:
        raise SystemExit("見出しが見つからない: " + head)
    st = hits[-1]
    # 付録の中にも「第1章　…」という小見出しが並ぶので、
    # 節の終わりは次の付録見出しだけで判定する。
    nxt = re.compile(r"^付録\s*[A-Z]" if not end_head else
                     "^(?:付録\s*[A-Z]|" + re.escape(end_head) + ")")
    for j in range(st + 1, len(ps)):
        if nxt.match(ps[j].text.strip()):
            return st, j
    return st, len(ps)


def main(docx, ledger, sheet, cites_json, head, note=None, end_head=None):
    doc = Document(docx)
    st, en = find_section(doc, head, end_head)
    print("付録 %s: 段落 %d〜%d（%d 段落を置き換える）" % (head, st, en, en - st - 1))
    tail = len(doc.paragraphs) - 1          # 末尾の段落番号（差し込み位置の目印）
    anchor = doc.paragraphs[st]._element
    for p in doc.paragraphs[st + 1:en]:     # 見出しは残し、中身だけ消す
        p._element.getparent().remove(p._element)
    doc.save(docx)

    meta = applycites.load_meta(ledger, sheet)
    cites = json.load(io.open(cites_json, encoding="utf-8"))
    order, unused = citedocx.apply(
        docx, cites, meta, out=docx,
        lead_text=applycites.LEAD % os.path.basename(ledger), extra_note=note)
    print("引用文献 %d 件 / 対応が外れたキー %d" % (len(order), len(unused)))

    # 末尾に生成された一覧を、付録の位置へ移す
    doc = Document(docx)
    st, _ = find_section(doc, head, end_head)
    body = doc.element.body
    kids = list(body.iterchildren())
    hi = kids.index(doc.paragraphs[st]._element)
    # 生成ブロック＝「参考文献」見出しからあと（改頁の段落を含む）
    gen = [i for i, p in enumerate(doc.paragraphs)
           if p.text.strip() == "参考文献" and i > st]
    if not gen:
        raise SystemExit("生成した一覧が見つからない")
    g0 = kids.index(doc.paragraphs[gen[-1]]._element)
    moved = [e for e in kids[g0:] if e.tag != W + "sectPr"]
    # 改頁だけの段落（見出しの直前）も一緒に持っていく
    if g0 > 0 and kids[g0 - 1].tag == W + "p" and not kids[g0 - 1].xpath("string(.)").strip():
        moved.insert(0, kids[g0 - 1])
    prev = doc.paragraphs[st]._element
    for e in moved:
        if e.xpath("string(.)").strip() == "参考文献":
            e.getparent().remove(e)         # 付録の見出しが既にある
            continue
        e.getparent().remove(e)
        prev.addnext(e)
        prev = e
    doc.save(docx)
    print("一覧を付録の位置へ移した（%d 要素）" % len(moved))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main(*sys.argv[1:])
