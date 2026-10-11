# -*- coding: utf-8 -*-
"""できあがった docx に、肩付き番号と巻末の参考文献一覧を後から入れる。

本ごとに生成スクリプトの作りが違うので、組版済みの docx を直接いじる。
やることは二つだけ。

  1. 本文の文を走査し、citesNN.py に載っている文の末尾（句点の直前）に
     肩付きの番号を差し込む
  2. 末尾に「参考文献」の節を足し、番号順に並べる

番号は本文に現れた順に 1 から振る。同じ文献が二度出てきたら同じ番号になる。
"""
import io
import os
import re
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SENT = re.compile(r"[^。]*。|[^。]+$")
NAVY = RGBColor(0x0A, 0x18, 0x30)
GREY = RGBColor(0x55, 0x60, 0x6E)


def split_sents(text):
    return [m.group(0) for m in SENT.finditer(text) if m.group(0)]


def h(sent):
    import hashlib
    return hashlib.md5(re.sub(r"\s+", "", sent).encode("utf-8")).hexdigest()[:10]


def fmt_range(nums):
    nums = sorted(set(nums))
    out = []
    i = 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j - i >= 2:
            out.append("%d-%d" % (nums[i], nums[j]))
        else:
            out.extend(str(n) for n in nums[i:j + 1])
        i = j + 1
    return ",".join(out)


def _copy_fmt(src, dst):
    if src is None:
        return
    dst.bold = src.bold
    dst.italic = src.italic
    if src.font.size is not None:
        dst.font.size = src.font.size
    if src.font.name:
        dst.font.name = src.font.name
        rpr = dst._element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), src.font.name)
    if src.font.color is not None and src.font.color.rgb is not None:
        dst.font.color.rgb = src.font.color.rgb


def insert_marks(par, marks):
    """marks: [(文字位置, 表示する番号)] を、段落の該当位置に肩付きで挿入する。

    段落は複数の run に分かれていることがあるので、位置を含む run を割って挟む。
    後ろから処理して、位置がずれないようにする。
    """
    for pos, label in sorted(marks, reverse=True):
        acc = 0
        target = None
        for r in par.runs:
            n = len(r.text)
            if acc <= pos <= acc + n:
                target = (r, pos - acc)
                break
            acc += n
        if target is None:
            continue
        run, off = target
        head, tail = run.text[:off], run.text[off:]
        run.text = head
        # 肩付きの run
        sup = par.add_run(label)
        _copy_fmt(run, sup)
        if run.font.size is not None:
            sup.font.size = Pt(run.font.size.pt * 0.92)
        sup.font.superscript = True
        run._element.addnext(sup._element)
        if tail:
            rest = par.add_run(tail)
            _copy_fmt(run, rest)
            sup._element.addnext(rest._element)


def apply(path, cites, meta, out=None, lead_text=None, extra_note=None,
          title="参考文献"):
    """cites: {文のハッシュ: [PMID]}, meta: {PMID: 書誌の辞書}"""
    doc = Document(path)
    order = []
    used = set()

    for par in doc.paragraphs + [p for t in doc.tables for row in t.rows
                                 for c in row.cells for p in c.paragraphs]:
        text = par.text
        if not text or "{" in text:
            pass
        marks = []
        pos = 0
        for s in split_sents(text):
            k = h(s)
            if k in cites:
                nums = []
                for pm in cites[k]:
                    if pm not in order:
                        order.append(pm)
                    nums.append(order.index(pm) + 1)
                # 句点の手前に置く
                at = pos + len(s) - (1 if s.endswith("。") else 0)
                marks.append((at, fmt_range(nums)))
                used.add(k)
            pos += len(s)
        if marks:
            insert_marks(par, marks)

    # ---- 参考文献
    doc.add_page_break()
    hp = doc.add_paragraph()
    hr = hp.add_run(title)
    hr.bold = True
    hr.font.size = Pt(14)
    hr.font.color.rgb = NAVY
    hp.paragraph_format.space_after = Pt(6)
    for t in ([lead_text] if lead_text else []) + ([extra_note] if extra_note else []):
        p = doc.add_paragraph()
        r = p.add_run(t)
        r.font.size = Pt(10.5)
        p.paragraph_format.line_spacing = 1.18
        p.paragraph_format.space_after = Pt(4)
    for i, pm in enumerate(order, 1):
        m = meta[pm]
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.14
        pf.left_indent = Cm(0.95)
        pf.first_line_indent = Cm(-0.95)
        r0 = p.add_run("%d) " % i)
        r0.bold = True
        r0.font.size = Pt(8.8)
        r0.font.color.rgb = NAVY
        au = m.get("first", "")
        # 日本語の発行者（学会・省庁）には「, et al.」を付けない
        if au and not re.search(r"[ぁ-んァ-ヶ一-龥]", au):
            au += ", et al."
        desc = "%s %s. %s. %s." % (au, (m.get("title") or "").rstrip("."),
                                   m.get("journal", ""), m.get("year", ""))
        r1 = p.add_run(desc.strip())
        r1.font.size = Pt(8.8)
        tail = []
        pmid = m.get("pmid", pm)
        if pmid and pmid.isdigit():
            tail.append("PMID %s" % pmid)
        if m.get("doi"):
            tail.append("doi:%s" % m["doi"])
        if tail:
            r2 = p.add_run("　" + "　".join(tail))
            r2.font.size = Pt(8.2)
            r2.font.color.rgb = GREY
    doc.save(out or path)
    return order, set(cites) - used


def strip_refs(doc, title="参考文献"):
    """末尾の参考文献の節（見出しから最後まで）を取り除く。作り直すときに使う。"""
    body = doc.element.body
    ps = list(body.iterchildren())
    start = None
    for i, el in enumerate(ps):
        texts = "".join(t.text or "" for t in el.iter(qn("w:t")))
        if texts.strip() == title:
            start = i
    if start is None:
        return 0
    # 直前の改頁も落とす
    j = start - 1
    while j >= 0:
        txt = "".join(t.text or "" for t in ps[j].iter(qn("w:t"))).strip()
        has_break = ps[j].findall(".//" + qn("w:br"))
        if not txt and has_break:
            start = j
            break
        if txt:
            break
        j -= 1
    n = 0
    for el in ps[start:]:
        # 版面の設定（sectPr）は消さない。消すと Word が既定の余白に戻し、
        # ページ数が大きく変わってしまう。
        if el.tag == qn("w:sectPr"):
            continue
        body.remove(el)
        n += 1
    return n


def redo_refs(path, cites, meta, out=None, lead_text=None, extra_note=None,
              title="参考文献"):
    """本文の肩付き番号はそのままに、末尾の一覧だけ作り直す。"""
    doc = Document(path)
    order = []
    for par in doc.paragraphs + [p for t in doc.tables for row in t.rows
                                 for c in row.cells for p in c.paragraphs]:
        for s in split_sents(par.text):
            k = h(re.sub(r"[0-9,\-]+$", "", s.rstrip("。")) + ("。" if s.endswith("。") else ""))
            if k in cites:
                for pm in cites[k]:
                    if pm not in order:
                        order.append(pm)
    strip_refs(doc, title)
    _write_refs(doc, order, meta, lead_text, extra_note, title)
    doc.save(out or path)
    return order


def _write_refs(doc, order, meta, lead_text, extra_note, title):
    doc.add_page_break()
    hp = doc.add_paragraph()
    hr = hp.add_run(title)
    hr.bold = True
    hr.font.size = Pt(14)
    hr.font.color.rgb = NAVY
    hp.paragraph_format.space_after = Pt(6)
    for t in ([lead_text] if lead_text else []) + ([extra_note] if extra_note else []):
        p = doc.add_paragraph()
        r = p.add_run(t)
        r.font.size = Pt(10.5)
        p.paragraph_format.line_spacing = 1.18
        p.paragraph_format.space_after = Pt(4)
    for i, pm in enumerate(order, 1):
        m = meta[pm]
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.14
        pf.left_indent = Cm(0.95)
        pf.first_line_indent = Cm(-0.95)
        r0 = p.add_run("%d) " % i)
        r0.bold = True
        r0.font.size = Pt(8.8)
        r0.font.color.rgb = NAVY
        au = m.get("first", "")
        if au and not re.search(r"[ぁ-んァ-ヶ一-龥]", au):
            au += ", et al."
        desc = "%s %s. %s. %s." % (au, (m.get("title") or "").rstrip("."),
                                   m.get("journal", ""), m.get("year", ""))
        r1 = p.add_run(desc.strip())
        r1.font.size = Pt(8.8)
        tail = []
        pmid = m.get("pmid", pm)
        if pmid and pmid.isdigit():
            tail.append("PMID %s" % pmid)
        if m.get("doi"):
            tail.append("doi:%s" % m["doi"])
        if tail:
            r2 = p.add_run("　" + "　".join(tail))
            r2.font.size = Pt(8.2)
            r2.font.color.rgb = GREY
