import json, re, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = 'https://www.j-circ.or.jp/cms/wp-content/uploads/'
G = {  # file: (title, year/edition, group, upload path)
 'JCS2022_hiraoka': ('非心臓手術における合併心疾患の評価と管理に関するガイドライン', '2022年改訂版', 'A 心疾患合併患者の非心臓手術', '2022/03'),
 'JCS2026_Kamiya_Katsuragi': ('心疾患患者の妊娠・出産の適応、管理に関するガイドライン', '2026年改訂版', 'A 心疾患合併患者の非心臓手術', '2026/03'),
 'JCS2025_Yamagishi': ('成人先天性心疾患診療ガイドライン', '2025年改訂版', 'A 心疾患合併患者の非心臓手術', '2025/03'),
 'JCS2025_Tamura': ('肺血栓塞栓症・深部静脈血栓症および肺高血圧症に関するガイドライン', '2025年改訂版', 'A 心疾患合併患者の非心臓手術', '2025/03'),
 'JCS2020_Kimura_Nakamura': ('冠動脈疾患患者における抗血栓療法', '2020年JCSガイドライン フォーカスアップデート版', 'A 心疾患合併患者の非心臓手術', '2020/04'),
 'JCS2024_Iwasaki': ('不整脈治療', '2024年JCS/JHRSフォーカスアップデート版', 'A 心疾患合併患者の非心臓手術', '2024/03'),
 'JCS2022_Ohuchi_Kawada': ('先天性心疾患術後遠隔期の管理・侵襲的治療に関するガイドライン', '2022年改訂版', 'A 心疾患合併患者の非心臓手術', '2022/03'),
 'JCS2025_Kato': ('心不全診療ガイドライン', '2025年改訂版', 'B 補助循環・心不全', '2025/03'),
 'JCS2021_Ono_Yamaguchi': ('重症心不全に対する植込型補助人工心臓治療ガイドライン', '2021年改訂版', 'B 補助循環・心不全', '2021/03'),
 'JCS2023_nishimura': ('PCPS/ECMO/循環補助用心内留置型ポンプカテーテルの適応・操作', '2023年JCS/JSCVS/JCC/CVITフォーカスアップデート版', 'B 補助循環・心不全', '2023/03'),
 'JCS2025_Saiki': ('心臓移植に関するガイドライン', '2025年改訂版', 'B 補助循環・心不全', '2025/03'),
 'JCS2018_kimura': ('急性冠症候群ガイドライン', '2018年改訂版', 'B 補助循環・心不全', '2018/11'),
 'JCS2020_Izumi_Eishi': ('弁膜症治療のガイドライン', '2020年改訂版', 'C Structural（カテーテル弁治療など）', '2020/04'),
 'JCS2021_Sakamoto_Kawamura': ('先天性心疾患、心臓大血管の構造的疾患（structural heart disease）に対するカテーテル治療のガイドライン', '2021年改訂版', 'C Structural（カテーテル弁治療など）', '2021/03'),
 'JCS2021_Kurita_Nogami': ('不整脈非薬物治療', '2021年JCS/JHRSフォーカスアップデート版', 'C Structural（カテーテル弁治療など）', '2021/03'),
 'JCS2026_Izumi': ('感染性心内膜炎診療ガイドライン', '2026年改訂版', 'C Structural（カテーテル弁治療など）', '2026/03'),
 'JCS2018_nakamura_yaku': ('安定冠動脈疾患の血行再建ガイドライン', '2018年改訂版', 'D その他（冠動脈・大血管）', '2018/09'),
 'JCS2020_Ogino': ('大動脈瘤・大動脈解離診療ガイドライン', '2020年改訂版', 'D その他（冠動脈・大血管）', '2020/07'),
 'JCS2025_Shigematsu': ('末梢動脈疾患ガイドライン', '2025年フォーカスアップデート版', 'D その他（冠動脈・大血管）', '2025/03'),
 'JCS2024_Tsukada_Tetsuo': ('多様性に配慮した循環器診療ガイドライン', '2024年', 'D その他（冠動脈・大血管）', '2024/03'),
}
KEYS = ['麻酔・鎮静', 'TEE', '周術期', '補助循環', 'Structural', '抗血栓の休薬']
res = json.load(open('sections.json', encoding='utf-8'))

wb = openpyxl.Workbook()
ws0 = wb.active; ws0.title = 'ガイドライン一覧'
ws0.append(['分類', 'ガイドライン', '版', 'PDF（クリックで開く）', '総頁', '該当節の数'] + KEYS)
ws = wb.create_sheet('該当節の索引')
ws.append(['分類', 'ガイドライン', '節（章 > 節 > 項）', '頁', 'PDFの該当頁'] + KEYS + ['本文の冒頭'])
for f, (title, ed, grp, path) in sorted(G.items(), key=lambda x: (x[1][2], x[0])):
    url = BASE + path + '/' + f + '.pdf'
    secs = res[f]
    hits = []
    for s in secs:
        c = s['counts']
        score = c['麻酔・鎮静'] * 2 + c['TEE'] * 2 + c['補助循環'] + c['Structural'] + c['抗血栓の休薬'] + c['周術期'] * 0.5
        if score >= 3 and s['len'] > 150:
            hits.append(s)
    npages = max((p for s in secs for p in s['pages']), default=0)
    tot = [sum(s['counts'][k] for s in hits) for k in KEYS]
    ws0.append([grp, title, ed, '開く', npages, len(hits)] + tot)
    c = ws0.cell(ws0.max_row, 4); c.hyperlink = url; c.font = Font(color='0563C1', underline='single')
    for s in hits:
        p0 = s['pages'][0]
        pr = f"{p0}" if p0 == s['pages'][-1] else f"{p0}–{s['pages'][-1]}"
        head = re.sub(r'\s+', ' ', s['section'])
        ws.append([grp, title + '（' + ed + '）', head, pr, f'{p0}頁を開く'] + [s['counts'][k] for k in KEYS] + [s['text'][:160]])
        c = ws.cell(ws.max_row, 5); c.hyperlink = f'{url}#page={p0}'; c.font = Font(color='0563C1', underline='single')

# non-JCS documents
ws2 = wb.create_sheet('適正使用指針など')
ws2.append(['対象', '文書', '作成団体', '版', 'URL', '麻酔・周術期に関する記載', '確認'])
rows = [
    ['Impella', 'IMPELLA 適正使用指針（細則を含む）', '補助人工心臓治療関連学会協議会 インペラ部会（J-PVAD）', '第6版 2024年6月18日改定', 'https://j-pvad.jp/guidance/',
     '適応（心原性ショック等の薬物療法抵抗性の急性心不全、IABP/PCPSのみでは不十分な病態。PCI中のみの補助には使わない）、施設要件（心臓血管手術年100例以上、植込型VAD施設との連携、体外循環技術認定士、補助循環に習熟したCCU/ICU）、全例レジストリ。麻酔・抗凝固・管理の記載はない', 'くろたんが原文で確認'],
    ['ASD・PFO閉鎖', '経皮的心房中隔欠損閉鎖セットに関する適正使用指針', '日本心血管インターベンション治療学会（CVIT）ほか', '2024年11月1日', 'http://www.cvit.jp/docs/about/committee/device/asd_pda/guideline-occluder-atrial_septum.pdf',
     '施設要件「全身麻酔ができ、リアルタイム3次元心エコー装置などを置く十分なスペースのあるカテーテル室、ハイブリッド手術室を有する」（Copilot の引用、原文未照合）', 'Copilot（原文未照合）'],
    ['TPVI', 'Sapien3／Harmony による経カテーテル肺動脈弁留置術の施設・術者認定', '経カテーテル的心臓弁治療関連学会協議会・日本小児循環器学会', '2024年9月／2025年3月 一部改訂', '',
     '「麻酔科医による全身麻酔管理が可能であること」「術中経食道心エコー検査ができる体制」（Copilot の引用、URL・原文未照合）', 'Copilot（原文未照合）'],
]
for r in rows:
    ws2.append(r)
    if r[4]:
        c = ws2.cell(ws2.max_row, 5); c.hyperlink = r[4]; c.font = Font(color='0563C1', underline='single')

def fmt(w, widths, hdr):
    for i, wd in enumerate(widths, 1):
        w.column_dimensions[get_column_letter(i)].width = wd
    for row in w.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    for c in w[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=hdr)
    w.freeze_panes = 'A2'; w.auto_filter.ref = w.dimensions
fmt(ws0, [26, 48, 22, 10, 7, 8] + [8] * 6, '1F4E79')
fmt(ws, [22, 40, 50, 8, 11] + [7] * 6 + [60], '375623')
fmt(ws2, [12, 40, 32, 22, 30, 70, 18], '7030A0')
wb.save('index.xlsx')
print(ws.max_row - 1, 'rows')
