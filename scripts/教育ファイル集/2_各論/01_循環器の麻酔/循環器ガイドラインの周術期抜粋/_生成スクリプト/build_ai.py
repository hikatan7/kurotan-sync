import docx, re
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import docx.opc.constants

FONT = 'メイリオ'
d = docx.Document()
sec = d.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = sec.bottom_margin = Cm(1.6)
st = d.styles['Normal']; st.font.name = FONT; st.font.size = Pt(9.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
st.paragraph_format.space_after = Pt(2); st.paragraph_format.line_spacing = 1.15
for name, size in [('Heading 1', 13.5), ('Heading 2', 11)]:
    h = d.styles[name]; h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True
    h.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79); h.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    h.paragraph_format.space_before = Pt(8); h.paragraph_format.space_after = Pt(3)


def hyperlink(p, url, text, size=8.5):
    part = p.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), r_id)
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '0563C1'); rPr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); rPr.append(sz)
    r.append(rPr); t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve'); r.append(t)
    h.append(r); p._p.append(h)


def fill(p, text, size):
    # text may contain <URL> markers -> hyperlinks
    for part in re.split(r'(<https?://[^>]+>)', text):
        if not part:
            continue
        if part.startswith('<http'):
            u = part[1:-1]; hyperlink(p, u, u, size)
        else:
            r = p.add_run(part); r.font.size = Pt(size)


def para(text, size=9.5):
    p = d.add_paragraph(); fill(p, text, size); return p


def bullets(items, size=9.5):
    for t in items:
        p = d.add_paragraph(style='List Bullet'); fill(p, t, size); p.paragraph_format.space_after = Pt(1)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), color); tcPr.append(s)


def table(header, rows, widths, hdr='1F4E79', size=8.5):
    t = d.add_table(rows=1, cols=len(header)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(h); r.bold = True
        r.font.size = Pt(size); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, hdr)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''; fill(cells[i].paragraphs[0], str(v), size)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    d.add_paragraph().paragraph_format.space_after = Pt(2)


p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('他AI横断検索の回収と照合'); r.bold = True; r.font.size = Pt(16)
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Structural・補助循環の国内文書／国際ガイドライン／TAVI・TEERの麻酔法／残っていた検索').font.size = Pt(10.5)
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p.add_run('作成：麻酔科　玉里順子　2026年10月12日').font.size = Pt(9)
para('2026年10月11〜12日に Gemini・ChatGPT・Copilot・Perplexity・Genspark に分担して投げた検索の回答を回収し、PubMed・各学会の原文で照合した結果。「照合」欄は、◎＝原文で一致、△＝一部誤りを訂正済み、×＝誤り、－＝原文未照合。')

# ---------------- 1
d.add_heading('1. Structural・補助循環の国内文書（施設基準・適正使用指針）', level=1)
table(['対象', '文書（版）と原文の要点', '照合'], [
    ['TAVI', 'THT協議会「経カテーテル的大動脈弁置換術 実施施設基準」（2013年7月18日施行、最終 2026年6月8日改定）。「麻酔科医（常勤*）が1名以上在籍すること」「施設として、麻酔科医/体外循環技術認定士の緊急動員に配慮すること」「経食道心エコー検査が年間200例以上」「術中経食道心エコー検査が実施可能」。<http://j-tavr.com/guideline.html>', '△ Geminiは改定年月を2025年11月と回答。原文は2026年6月8日'],
    ['MitraClip・PASCAL', '日本循環器学会ほか「弁尖間クリッピング式の経皮的僧帽弁接合不全修復システムに関する適正使用指針」（2017年11月18日初版、2024年12月6日改訂）。「全身麻酔ができ、3次元心エコー装置などを置く十分なスペースのあるカテーテル室を有する。ハイブリッド手術室が推奨される」「経食道心エコー検査を年間100例以上（うち術中30例以上、診断目的50例以上）」。手術リスクは循環器内科医と心臓血管外科医のハートチームで評価。麻酔科医の明記はない。<https://www.j-circ.or.jp/cms/wp-content/uploads/2024/12/MTEERusage-guidelines20241206.pdf>', '◎'],
    ['TriClip', '「弁尖間クリッピング式の経皮的三尖弁接合不全修復システムに関する適正使用指針」（2025年9月10日初版）。ハートチームによるリスク評価のみ。麻酔・TEEの記載はない。<https://www.j-circ.or.jp/cms/wp-content/uploads/2025/12/TriClipusage-guidelines20250910_20251208.pdf>', '◎'],
    ['WATCHMAN', '「左心耳閉鎖システムに関する適正使用指針」（2024年12月6日改訂）。「全身麻酔ができ、経食道心エコーや麻酔などの装置を置く十分なスペースがあるカテーテル室を有する。ハイブリッド手術室を有することが望ましい」「術中経食道心エコー検査が実施可能であること」「心臓外科医による緊急開胸手術が遅滞なく実施可能」。<https://www.j-circ.or.jp/cms/wp-content/uploads/2024/12/WMusage-guidelines20241206.pdf>', '◎'],
    ['ASD・PFO閉鎖', 'CVITほか「経皮的心房中隔欠損閉鎖セットに関する適正使用指針」（2024年11月1日作成、2025年12月1日改訂）。「全身麻酔ができ、リアルタイム3次元心エコー装置などを置く十分なスペースのあるカテーテル室、ハイブリッド手術室を有する」「経食道心エコー検査を年間100例以上」。<http://www.cvit.jp/docs/about/committee/device/asd_pda/guideline-occluder-atrial_septum.pdf>', '◎（Copilotは改訂日を欠落）'],
    ['TPVI', 'THT協議会・日本小児循環器学会「Sapien3／Harmony による経カテーテル肺動脈弁留置術の施設・術者認定」。「麻酔科医による全身麻酔管理が可能であること」「術中経食道心エコー検査ができる体制」', '－ URL未取得'],
    ['IMPELLA', '補助人工心臓治療関連学会協議会インペラ部会「IMPELLA適正使用指針」第6版（2024年6月18日改定）。適応・施設要件・全例登録のみで、麻酔・抗凝固の記載はない。<https://j-pvad.jp/guidance/>', '◎'],
    ['麻酔科学会系', '日本麻酔科学会・日本心臓血管麻酔学会による TAVI・MitraClip・左心耳閉鎖の麻酔管理指針は、見つからなかった（Copilot）', '－'],
    ['診療報酬', 'K555-2・K559-3・K594-2・K602-2 の施設基準の麻酔科要件は、令和8年3月5日保医発0305第8号の別添を直接読む必要がある（Copilotは原文に到達できず）', '－ 未取得'],
], [2.4, 11.6, 3.0], hdr='7030A0')

# ---------------- 2
d.add_heading('2. 国際ガイドライン（Perplexity）', level=1)
table(['文書', '取れた記載', '照合'], [
    ['2025 ESC/EACTS 弁膜症ガイドライン。Eur Heart J 2025;46(44):4635–4736. PMID 40878295', '治療手段はハートチームが臨床・解剖・手技の特徴と生涯管理を踏まえて決める。OAC適応のない患者のTAVI後は低用量アスピリン（DOACのルーチン使用、DAPTは推奨しない）。M-TEERは二次性MRで推奨（Class I・LOE A と総説で要約）。経カテーテル三尖弁治療は症候性高リスクで Class IIa。TAVIの局所麻酔／全身麻酔、TEERのTEEガイドに関する個別推奨は、取得範囲では見つからなかった', '△ PerplexityのPMID 42064867 は別の総説（Adamo M, Eur Heart J Suppl 2026）。正しくは40878295'],
    ['2020 ACC/AHA 弁膜症ガイドライン。Circulation 2021;143(5):e72–e227. PMID 33332150', 'ハートチームの構成員として「cardiovascular anesthesiologists」を挙げる（推奨本文ではなく解説）。弁膜症センターの設備として cardiac anesthesia support、IABP・経皮的VAD・ECMO', '◎ PMIDはPerplexity未記載、DOIから特定'],
    ['補助循環の専門家合意（SCAI・ACC・ESC-ACVC）', 'DanGer Shock（NEJM 2024;390:1382–1393）以後の Impella・IABP の抗凝固・アクセス・離脱に関する逐語推奨は取れなかった', '－'],
], [5.0, 9.0, 3.0], hdr='375623')

# ---------------- 3
d.add_heading('3. TAVI・TEER の麻酔法の研究（ChatGPT、書誌は PubMed で訂正済み）', level=1)
table(['論文', '要点（抄録の数値）', '照合'], [
    ['Thiele H, et al. Circulation 2020;142(15):1437–1447. PMID 32819145（SOLVE-TAVI）', 'RCT 447例、意識下鎮静（CS）対 全身麻酔（GA）。30日の複合エンドポイント 27.2% 対 26.4%（同等性 P=0.015）', '◎'],
    ['Butala NM, et al. JACC Cardiovasc Interv 2020;13(11):1277–1287. PMID 32499018', 'STS/ACC TVT 経大腿 TAVR 120,080例・559施設、操作変数法。CS で院内死亡（調整差 0.2%、p=0.010）・30日死亡（0.5%、p<0.001）が低く、在院日数が0.8日短い', '△ ChatGPTは巻号頁を 13(18):2117–2128 と誤記'],
    ['Hyman MC, et al. Circulation 2017;136(22):2132–2140. PMID 28864443', 'TVT Registry 10,997例（CS 1,737例）。CSで院内死亡 1.6% 対 2.5%、30日死亡 2.9% 対 4.1%（いずれも P=0.03）。CSからGAへの移行 5.9%', '△ ChatGPTは誌名を J Am Coll Cardiol 69(17) と誤記、DOIも誤り'],
    ['Maas EH, et al. Curr Pharm Des 2016;22(13):1868–1878. PMID 26642777', 'TAVIの全身麻酔対局所麻酔のシステマティックレビュー・メタ解析（統合研究数・効果量は本文で要確認）', '◎ 書誌のみ'],
    ['Horn P, et al. Catheter Cardiovasc Interv 2017;90(7):1212–1219. PMID 28112459', 'MitraClip 232例、深鎮静 対 全身麻酔。安全性の複合エンドポイントに差なし。深鎮静から全身麻酔への移行 2%。準備時間・ICU滞在は深鎮静で短い', '△ ChatGPTは巻号頁を 89(6):E163–E168 と誤記'],
    ['Ledwoch J, et al. Clin Res Cardiol 2016;105(4):297–306. PMID 26377429', 'MitraClip 連続91例（全身麻酔28・鎮静あり局所35・鎮静なし局所28）。手技成功 89%で群間差なし。ICU/IMC 入室は 100%・14%・14%', '△ ChatGPTは筆頭著者・誌名を空欄'],
], [5.6, 8.6, 2.8], hdr='1F4E79')

# ---------------- 4
d.add_heading('4. 残っていた時間のかかる検索', level=1)
table(['テーマ', '結果', '照合'], [
    ['JRC蘇生ガイドライン2025 の開心術後心停止（Gemini）', '心臓手術後の心停止（CALS）に関する章・CQは見つからない、との回答。特殊な状況の心停止として溺水・肺血栓塞栓症・オピオイド中毒、腹臥位CPRへの言及。CALS Japan は特定非営利活動法人 <https://www.calsjapan.org/>', '－ 章立ては二次資料からの再構成で未照合。CALS Japan のサイトは実在'],
    ['AHA 2025 Part 12 蘇生教育（Perplexity）', 'scripted debriefing COR 2b・LOE C-LD（新規）。cognitive aid は医療者 2a、市民 3:No benefit。RCDP 2b。spaced learning 2a・B-R。massed + booster 1・B-R。feedback device は医療者・市民とも 1・A。<https://cpr.heart.org/en/resuscitation-science/cpr-and-ecc-guidelines/resuscitation-education-science>', '－ 推奨文はPerplexityの引用'],
    ['Andreatta 2011 の「33%→50%」', 'Pediatr Crit Care Med 2011;12(1):33–38. PMID 20581734。抄録は「approximately 50%」のみで、33%は二次資料の記載', '－'],
    ['AKI 回復期の尿細管機能（ChatGPT）', '各機能の回復週数を定量した研究は見つからない。関連：Hong 1979 Kidney Int（PMID 390218、移植後ATNでFENaは回復とともに正常化）、Abdulkader 1992 Ren Fail（PMID 1561389、ATN後2〜36か月で43%にGFR低下と酸排泄低下）、Moran 1985 JCI（PMID 2414321）', '－'],
    ['成人先天性心疾患の診療体制（Genspark）', '日本成人先天性心疾患学会の専門医制度は2019年4月1日開始。専門医数・修練施設数は一覧が動的ページで取れず。移行期医療支援センターは各都道府県1か所以上の設置を求める（厚労省ガイド、平成29年）。「ACHD患者50万人以上、年1万人増」は JCS「先天性心疾患の成人への移行医療に関する提言（2022年第3版）」の記載で、原典論文は特定できず。2025年版ACHDガイドラインは2026年5月22日更新', '－'],
    ['静岡県の二次医療圏別 肺がん罹患数（Copilot）', '県「全国がん登録」ページの年次確定時集計表（zip、2016〜2020年）に二次医療圏別・部位別がある。<https://www.pref.shizuoka.jp/kousei/ko-13/shippei/cancer/1003080.html>', '－ zip未展開'],
    ['心臓・肺移植の体制（Copilot）', 'JOTNW 施設一覧：心臓13・肺12施設。厚労省報告書（令和8年6月9日）：令和7年度の脳死下提供155名、心臓126件・肺149件。<https://www.mhlw.go.jp/content/001709253.pdf>', '－'],
], [3.8, 10.4, 2.8], hdr='C55A11')

# ---------------- 5
d.add_heading('5. 各AIの成績（この2日間）', level=1)
bullets([
    'Gemini：国内の学会文書を探す仕事は最も役に立った（JCS の M-TEER・TriClip・WATCHMAN 指針の URL と原文はすべて一致）。ただし TAVR 施設基準の改定年月を誤り、JRC2025 の章立ては二次資料から作った。送信ボタンが反応しないことが続き、タブが固まったことが1回あった。',
    'ChatGPT：実在する論文だけを挙げ、水増しはしない。ただし巻号頁・誌名・DOI を作る（6件中3件で誤り）。書誌は必ず PubMed から写す。',
    'Copilot：公的文書の探索で、見つからないものを「見つからない」と答えるようになった。一方で「想定」として推測の文書名を表に混ぜた回があった。統計局の URL を誤った。',
    'Perplexity：ガイドラインの本文引用は正確だが、PMID を別論文のものにしていた。本文の後半まで読めず、推奨表の Class・LOE を取れないことが多い。',
    'Genspark（AI Chat・Haiku 4.5）：1回目は検索せずに「見つからない」と返すことがあった。促すと検索するが、動的ページ（裁判所・学会の一覧）は読めない。',
])
d.save('ai_results.docx'); print('ok')
