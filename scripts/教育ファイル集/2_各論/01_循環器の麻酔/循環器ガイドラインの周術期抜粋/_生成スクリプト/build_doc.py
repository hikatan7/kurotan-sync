import docx, re
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'メイリオ'
d = docx.Document()
sec = d.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = sec.bottom_margin = Cm(1.6)

st = d.styles['Normal']
st.font.name = FONT; st.font.size = Pt(9.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
st.paragraph_format.space_after = Pt(2)
st.paragraph_format.line_spacing = 1.15
for name, size in [('Heading 1', 14), ('Heading 2', 11.5), ('Heading 3', 10)]:
    h = d.styles[name]; h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True
    h.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    h.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    h.paragraph_format.space_before = Pt(8 if name != 'Heading 3' else 5); h.paragraph_format.space_after = Pt(3)


def add_runs(p, text):
    # [n] -> superscript citation
    for part in re.split(r'(\[\d+(?:[,–]\d+)*\])', text):
        if not part:
            continue
        m = re.fullmatch(r'\[(.+)\]', part)
        r = p.add_run(m.group(1) if m else part)
        if m:
            r.font.superscript = True
    return p


def para(text, bold_lead=None, size=None):
    p = d.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True
    add_runs(p, text)
    if size:
        for r in p.runs: r.font.size = Pt(size)
    return p


def bullets(items):
    for t in items:
        p = d.add_paragraph(style='List Bullet')
        add_runs(p, t)
        p.paragraph_format.space_after = Pt(1)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), color); tcPr.append(s)


def table(header, rows, widths, hdr='1F4E79', size=8.5):
    t = d.add_table(rows=1, cols=len(header)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(h); r.bold = True
        r.font.size = Pt(size); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, hdr)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''; p = cells[i].paragraphs[0]; add_runs(p, str(v))
            for r in p.runs: r.font.size = Pt(size)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ---------------- title
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('循環器系ガイドラインにみる周術期管理'); r.bold = True; r.font.size = Pt(17)
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Structural Heart・補助循環・心疾患合併患者の非心臓手術 ── 麻酔科医のための抜粋'); r.font.size = Pt(11)
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p.add_run('作成：麻酔科　玉里順子').font.size = Pt(9.5)
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p.add_run('2026年10月11日').font.size = Pt(9)

para('日本循環器学会（JCS）と関連学会の合同ガイドライン、および補助循環・カテーテル治療の適正使用指針から、麻酔科医が周術期に判断を求められる記載を抜き出した。推奨クラス（I／IIa／IIb／III）とエビデンスレベル（A／B／C）は原典の表記のまま示し、頁は各ガイドラインPDFの頁である。全節の頁の索引は別添の「索引xlsx」にある。')

# ================= Part 1
d.add_heading('第1部　心疾患を合併する患者の非心臓手術', level=1)
para('2022年改訂版 非心臓手術における合併心疾患の評価と管理に関するガイドライン[1]が中心になる。Part 1が背景知識にもとづく推奨（第1〜5章）、Part 2がGRADEによるCQ1〜6（第6章）である。')

d.add_heading('1. 術前評価の流れ（第2章11、p.33–35）', level=2)
bullets([
    'Step 1 緊急手術か：緊急なら循環器精査で手術を遅らせない。病歴・身体所見・心電図で循環器緊急症の有無だけ把握する[1]。',
    'Step 2 循環器緊急症（ACS、重症不整脈、急性心不全、症候性弁膜症）があれば安定化を優先し、侵襲的治療を先行させるかは「心疾患の緊急度・非心臓手術の緊急度・両者のリスク・抗血栓薬の継続と中止」を個別に吟味する[1]。',
    'Step 3 低リスク手術なら手術へ。Step 4 RCRIで評価し、非血管手術は0〜1項目、血管手術は0項目なら手術へ。上昇リスク群のみStep 5（運動耐容能・BNP など）へ進む[1]。',
])
table(['推奨', 'クラス', 'Lv', '頁'], [
    ['術前リスク評価にはRCRIを使用する', 'I', 'B', '30'],
    ['RCRI低リスク（非血管0〜1、血管0項目）ではルーチンの追加精査をしない', 'III', 'B', '30'],
    ['自己申告の運動耐容能が7 METs以上なら追加精査をしないことを考慮', 'IIa', 'C', '26'],
    ['同 4 METs以上なら追加精査をしないことを考慮してもよい（DASIスコアの利用は IIa）', 'IIb', 'C', '26'],
    ['中・高リスク手術で上昇リスク群ではBNP/NT-proBNP測定を考慮してもよい', 'IIb', 'C', '30'],
    ['低リスク手術のルーチン術前12誘導心電図', 'III', 'B', '27'],
    ['ルーチンの術前経胸壁心エコー', 'III', 'C', '27'],
    ['ルーチンの冠動脈CT・冠動脈造影', 'III', 'C', '33'],
], [11.5, 1.5, 1.2, 1.2])

d.add_heading('2. 常用薬と抗血栓薬（第3章、p.35–43）', level=2)
table(['薬剤', '推奨', 'クラス', 'Lv', '頁'], [
    ['β遮断薬', '手術24時間以内に新規開始しない／長期服用中は継続', 'III／I', 'A／B', '35'],
    ['スタチン', '血管手術の2週間以上（できれば1ヵ月）前に開始を考慮／長期服用中は継続', 'IIa／I', 'B／C', '36'],
    ['ACE阻害薬・ARB', '手術当日の休薬を考慮してもよい／HFrEFで中止したら術後早期に再開', 'IIb／I', 'C', '36'],
    ['α2作動薬', '術前の新規開始は推奨されない', 'III', 'B', '36'],
    ['アスピリン', '一次予防目的の術前新規開始はしない／維持療法の継続・中止は血栓と出血のリスクで決める', 'III／I', 'B', '37'],
    ['抗凝固薬', '出血リスク低い手術は継続、中等度は個別判断、高い手術は休薬', 'I', 'A／C', '40'],
    ['ヘパリン置換', 'ワルファリン休薬時のルーチン置換はしない。機械弁・弁膜症性心房細動、塞栓高リスクかつ出血低リスクでは考慮（IIa）。DOAC休薬時の置換はしない', 'III／IIa／III', 'C／C／B', '42'],
    ['区域麻酔', '抗血栓薬服用中に区域麻酔を考える場合は、継続・中断を麻酔科医を含む集学的チームで協議', 'I', 'C', '43'],
], [2.6, 9.0, 1.7, 1.4, 1.0])
para('ワルファリンは5〜3日前に中止し、前日にINRを測る。<1.5なら予定どおり、1.5〜1.8ならリバースを考慮し、術日夕方か翌朝から維持量で再開する（表33、p.40）[1]。', size=9)

d.add_heading('3. 疾患別のマネージメント（第4章、p.44–69）', level=2)
d.add_heading('冠動脈疾患・ステント留置後', level=3)
bullets([
    '安定冠動脈疾患では、原則として術前に予防的な冠血行再建をしない（CQ1、GRADE 1B、p.72）[1]。',
    'ステント留置後1〜1.5ヵ月以内の待機的非心臓手術は行わない（表42 III-Harm／CQ2 GRADE 1C、p.50・77）。DAPT継続下が難しい手術は、DAPTが必要な期間が終わるまで延期を考慮する（IIa）[1]。',
    'ステント留置後で抗血小板薬服用中の患者は、出血リスクが高くなければアスピリン継続下で手術を行う（CQ3、GRADE 2C、p.80）[1]。',
])
d.add_heading('心不全・肥大型心筋症', level=3)
bullets([
    '新たに診断したHFrEFで3ヵ月以上延期できる中・高リスク手術は、GDMTを導入して心機能の回復を待つ（I-C）。服用中のβ遮断薬は周術期に中止しない（III-Harm B）（表47、p.56）[1]。',
    '肥大型心筋症：麻酔導入・陽圧換気による前負荷減少と末梢血管抵抗低下で左室流出路閉塞と低血圧が起こりうる。β遮断薬・非ジヒドロピリジン系Ca拮抗薬・シベンゾリンなどは継続し、前負荷を十分に保つ。流出路圧較差（安静時30 mmHg以上、Valsalva時50 mmHg以上）、SAMによる中等度以上のMR、LVEF<50%、拘束型の拡張障害は高リスク所見（表48、p.56–57）[1]。',
])
d.add_heading('肺高血圧症（p.57–61）', level=3)
bullets([
    '周術期には右心不全・低酸素・ショックを呈する肺高血圧クライシスが起こりやすい。体血圧が下がると右冠灌流が落ち、右室虚血から不可逆的な循環破綻に至る[1]。',
    '術前：麻酔科医・外科医・肺高血圧管理医の多分野チームで手術の必要性と管理を協議し、肺血管拡張薬は手術当日も中止しない。専門医不在なら専門施設への移送も考える[1]。',
    '術後対応の5段階（図12）：①増悪因子（低酸素、高CO2、アシドーシス、陽圧換気、内服中断、不整脈など）の是正 → ②右室前負荷の適正化（右房圧8〜12 mmHg目安、うっ血例は利尿・除水）→ ③体血圧の維持（ノルアドレナリン、バソプレシン）→ ④強心薬（ドブタミン、ミルリノン）→ ⑤肺血管拡張薬（NO吸入、PGI2吸入）[1]。',
    'リスク因子：NYHA III/IV、6MWD<300 m、RVSP/SBP≧0.66、右房圧>7 mmHg、緊急手術、中〜高リスク手術、麻酔3時間以上、術中昇圧薬使用、術中NO吸入の不使用など（表50）[1]。',
])
d.add_heading('弁膜症・高血圧・肺動脈カテーテル（p.61–64）', level=3)
bullets([
    '無症候性重症ASは、適切な血行動態モニター下での非心臓手術を考慮（IIa-B）。無症候性で左室機能の保たれた重症AR・MRも同様（IIa-C）（表52）[1]。',
    '無症候性重症ASでは予防的なSAVR・TAVIを行わずに非心臓手術をする（CQ4、GRADE 2D）。重症ASを合併する高齢者の大腿骨近位部骨折は、血行動態が安定していればAVRを先行させず、麻酔科・循環器内科・整形外科の集学的チームとHCU/ICUでの管理のもとで骨折手術を行う（CQ5、GRADE 2D、p.96）[1]。',
    'AS・MSでは頻脈を避けて前負荷を保ち、血管拡張薬は慎重に。AR・MRでは徐脈を避け、後負荷を下げる。高リスク例では術中に肺動脈カテーテルやTEEを検討する（p.61）[1]。',
    'I〜II度高血圧（<180/110 mmHg）なら手術を延期しないことを考慮してもよい（IIb-B）。周術期の血圧はベースラインの70〜100%に保つ（IIa-B）（表53）。非心臓手術での肺動脈カテーテルのルーチン使用は推奨されない（III-A）（表54）[1]。',
    'ワルファリン服用中の心房細動で CHADS2スコア4点以下なら、術前のヘパリン置換を行わない（CQ6、GRADE 2B、p.105）。5点以上では出血リスクが高くなければ置換を考慮する[1]。',
])
d.add_heading('植込み型心臓電気デバイス（CIED、p.67–69）', level=3)
bullets([
    '術前にデバイス種類、最終チェック日、自己脈の有無、ペーシングモード、ICD治療設定などを確認する（表57のチェック項目）[1]。',
    '自己脈がある症例は、電気メスの使用を数秒にとどめればモード変更は不要。非同期モードはspike on Tで心室性不整脈を誘発しうる[1]。',
    'ペースメーカは磁石でVOOになる。ICDは磁石で頻拍の感知がオフになるが、ペーシングモードは変わらない。非同期モードにはプログラマーが要る[1]。',
    'ICD・CRT-Dは除細動機能を手術直前にオフにし、体外式除細動パッドを前後（前胸部と背部）に、本体からできるだけ離して貼る。電気メスはバイポーラを使う[1]。',
])
d.add_heading('4. 術後心筋傷害（MINS、第5章、p.70–71）', level=2)
para('MINSは非心臓手術の約8〜18%に起こり、84〜93%は無症状である。VISION研究で30日死亡のHRは3.87。高リスク患者では術後の心筋トロポニン測定を考慮してもよいが、見つけたMINSへの治療は確立していない。発症時は循環器内科に相談する[1]。')

# ================= Part 2
d.add_heading('第2部　背景疾患ごとの周術期', level=1)
d.add_heading('1. 心疾患合併妊娠（2026年フォーカスアップデート版）', level=2)
para('国内の心疾患合併妊娠は全分娩の約1%。日産婦データベースでは心疾患群で帝王切開時の全身麻酔使用率（6.5% 対 4.2%）、無痛分娩率（4.0% 対 2.2%）が高く、母体死亡率も高かった（p.14–16）[6]。')
table(['対象', '推奨', '強度', 'Lv'], [
    ['PAH', '分娩方法は麻酔・PAHの循環管理体制を備えた施設で症例ごとに選ぶ（CQ8）', '弱い', 'C-LD'],
    ['PAH', '循環動態が安定したPAHの帝王切開は脊髄幹麻酔を推奨（CQ11）', '弱い', 'C-LD'],
    ['Fontan術後', '分娩方法はFontan循環の麻酔・循環管理体制を備えた施設で選ぶ（CQ10）', '弱い', 'C-LD'],
    ['Fontan術後', '帝王切開は脊髄幹麻酔を推奨（CQ12）', '弱い', 'C-LD'],
    ['Marfan症候群', '大動脈径45 mm超は帝王切開を推奨。40〜45 mmは一律の推奨なし（CQ9）', '弱い', 'C-LD'],
], [2.6, 10.6, 1.4, 1.4])
d.add_heading('2. 成人先天性心疾患の非心臓手術（2025年改訂版 第12・13章、p.72–75）', level=2)
bullets([
    '中等症以上では周術期の罹病率・死亡率・再挿管率が高い。先天性心疾患に習熟した麻酔科・循環器内科・心臓血管外科と看護師を含む多職種で当たる（I-C）[7]。',
    '高リスク病変：重度肺高血圧・Eisenmenger症候群、チアノーゼ型、NYHA III/IVの心室機能不全、左心系の高度狭窄、Fontan術後（表29）[7]。',
    'Eisenmenger症候群では、麻酔導入時に約1/4で低血圧、15〜20%で低酸素になる。導入時の血管収縮薬で低血圧を防げる[7]。',
    'Fontan術後のDOACは、通常リスク手術で半減期の2〜3倍、高出血リスクで4〜5倍の時間を休薬する。脳神経・眼科手術ではアスピリンを7日以上中止（IIb-B）[7]。',
    '麻酔法の優劣を示すエビデンスはない。術前の循環血液量と肺・体血管抵抗を維持するように調節する。疾患別の目標は表30（ASD・VSD・大動脈縮窄・Fallot四徴・修正大血管転位・TGA・単心室）。Williams症候群は全身麻酔・鎮静と心停止の関連に注意[7]。',
])
d.add_heading('3. 植込型LVAD装着時の麻酔（2021年改訂版、p.35–36）', level=2)
para('LVAD装着を受ける患者は収縮期血圧80 mmHg前後の低心拍出状態が多い。導入前に橈骨動脈ラインを確保し、手掌の末梢温、前額の脳酸素飽和度、麻酔深度モニターを付ける。少量の鎮静薬でも血圧が大きく下がる。左室機能が高度に不良な症例では、導入後に開胸前から大腿動静脈で人工心肺を始める。術中TEEでは右室機能、三尖弁逆流（中等度以上かつ弁輪径40 mm以上なら弁輪形成）、PFOなどの右左シャントを確認する[8]。')

# ================= Part 3
d.add_heading('第3部　補助循環（IABP・V-A ECMO・IMPELLA）', level=1)
para('2023年JCS/JSCVS/JCC/CVITフォーカスアップデート版「PCPS/ECMO/循環補助用心内留置型ポンプカテーテルの適応・操作」[2]と、IMPELLA適正使用指針（第6版）[9]から。')
table(['デバイス', '抗凝固の目安', '管理の要点', '頁'], [
    ['IABP', '駆動中は適切な抗凝固', '7〜8 Fr（下肢虚血が懸念されれば6 Fr、効果は減弱）。透視下で留置し、腹腔動脈にかからない長さを選ぶ。拡張は大動脈弁閉鎖時、収縮は拡張末期圧が最低になるよう調節。離脱は補助比率1:1→1:2→1:3と段階的に下げる。AR・大動脈瘤・解離・高度粥状硬化では使いにくい。IABP-SHOCK IIで死亡は改善せず（39.7% 対 41.3%）、ルーチン使用は推奨されない', '7, 18'],
    ['V-A ECMO', 'ACT 180〜200秒、APTT 50〜60秒（出血リスクがあればACT 160〜180秒）', '大腿送血で下肢虚血17%、切断5%。動脈圧波形・心エコーで大動脈弁の開放を確認。右上肢の血液ガス・SpO2で自己肺の酸素化を評価（north-south syndrome）。左室後負荷増加による肺うっ血・血栓に注意し、減圧はIABP併用・少量の強心薬・IMPELLA（ECPELLA）など', '19–20, 28'],
    ['IMPELLA', 'ACT 160〜180秒。パージ液はヘパリン50単位/mL（添付文書）。多くは全身ヘパリンの追加が要る', 'ポンプ・シャフトの血栓は塞栓と高度溶血の原因。離脱はPレベルを1〜2ずつ下げ、P2で安定を確かめてから心室外へ抜去し、抜去後も右心カテーテルで監視。国内では心原性ショック等、IABP/PCPSだけでは不十分な病態が適応で、PCI中のみの補助には使わない', '22–24'],
    ['ECPELLA', 'V-A ECMOに準じる（未分画ヘパリン）', 'V-A ECMO単独より溶血・出血は多いが脳卒中は差がない報告が多い。右上肢SpO2を指標にECMO・IMPELLA・人工呼吸の設定を調節。J-PVADの30日生存はIMPELLA単独81.1%、ECPELLA 49.6%', '7, 24–25'],
], [2.0, 3.6, 9.4, 1.2], hdr='375623')
bullets([
    'V-A ECMO中の強心薬は少量にとどめる（ドブタミン・ドパミン<3 μg/kg/分、ミルリノン<0.3 μg/kg/分）。ノルアドレナリンは<0.06 μg/kg/分程度にとどめ、前負荷と流量を確かめてから慎重に使う（p.28）[2]。',
    '鎮静：カニューラのトラブルは致命的なので安全を優先しつつ、不必要な深鎮静を避ける。RASS −2〜0を目標とすることが多い。不穏の原因（痛み、せん妄、低灌流、低酸素、感染）に対処する（p.29–30）[2]。',
    'IMPELLA実施施設の要件：心臓血管手術年100例以上、植込型VAD実施施設との連携、体外循環技術認定士、補助循環に習熟したCCU/ICU、全例レジストリ登録（第6版、2024年6月18日改定）[9]。',
])

# ================= Part 4
d.add_heading('第4部　Structural Heart Disease のカテーテル治療', level=1)
para('2021年改訂版 先天性心疾患、心臓大血管の構造的疾患に対するカテーテル治療のガイドライン[3]を中心に、弁膜症治療のガイドライン（2020年改訂版）[4]、不整脈非薬物治療（2021年フォーカスアップデート版）[5]、CVITと日本循環器学会ほかの適正使用指針、経カテーテル的心臓弁治療関連学会協議会の施設基準[10–14]から。')
table(['治療', '麻酔・画像・体制', '頁'], [
    ['TAVI', '術中TEEは術野と独立したエコー専門医が必要。局所麻酔下TAVIは全身麻酔と非劣性ないし優位の報告があるが、局所麻酔でも患者病態の変化に対応できるよう麻酔科管理は有用。実施施設基準：常勤麻酔科医1名以上、TEE年200例以上、ハイブリッド手術室（術中TEE、PCPS、緊急開心が可能）、麻酔科医・体外循環技術認定士の緊急動員への配慮（施設基準の最終改定は2026年6月8日。「麻酔科医（常勤）が1名以上在籍」は現行版でも同じ）', '[3] 72–73／[11]'],
    ['TAVI（アクセス）', 'アクセス評価は「麻酔リスクを判断する麻酔科医師」を含むHeart teamで行う（I-A）。全身麻酔が困難な患者に経心尖・直接大動脈アプローチを行うべきでない（III-Harm C）', '[3] 74'],
    ['PTAV（経皮的大動脈弁形成術）', '局所麻酔＋鎮静で行い、TEEで評価・モニターするなら全身麻酔。全身麻酔下の非心臓手術を控えた症候性重症ASで、術中の血行動態安定のためにPTAVを考慮してもよい（IIb-C）', '[3] 82–83'],
    ['MitraClip・PASCAL（経皮的僧帽弁接合不全修復）', '全身麻酔下・TEEガイド下に、経大腿静脈で心房中隔を穿刺して行う。TEEは適応決定・術中モニタリングとも「明らかな有用性」（3D TEEを含む）。適正使用指針（2024年12月6日改訂）の施設要件は、全身麻酔ができ3次元心エコー装置などを置く十分なスペースのあるカテーテル室（ハイブリッド手術室を推奨）、TEE年100例以上（うち術中30例以上）。手術リスクは循環器内科医と心臓血管外科医のハートチームで評価する', '[3] 18, 89／[12]'],
    ['TriClip（経皮的三尖弁接合不全修復）', '適正使用指針（2025年9月10日初版）は、外科手術が最適でないことをハートチームで評価することを求めるが、麻酔法・TEE・麻酔科の体制には触れていない', '[13]'],
    ['左心耳閉鎖（WATCHMAN）', '長期抗凝固療法が不適切な心房細動に考慮してもよい（IIb-B）。適正使用指針（2024年12月6日改訂）の施設要件は、全身麻酔ができTEEや麻酔の装置を置ける十分なスペースのあるカテーテル室（ハイブリッド手術室が望ましい）、術中TEEが実施可能、心臓外科医による緊急開胸が遅滞なく可能。術後はデバイス血栓をTEE・CTで経時的に評価', '[3] 18, 72／[5] 28–30／[14]'],
    ['ASD閉鎖', '施設基準に「心臓カテーテル室で麻酔科の管理の下に全身麻酔が可能」。CVIT指針でも「全身麻酔ができ、3次元心エコー装置を置けるカテーテル室・ハイブリッド手術室」、TEE年100例以上', '[3] 23／[10]'],
    ['PFO閉鎖', '全身麻酔ができTEEかICEが可能なカテーテル室。習熟までは全身麻酔下のTEEで行うことを推奨。診断時のTEEで鎮静を用いるとValsalva負荷が不十分になり、右左シャントを見逃しうる', '[3] 64–70'],
], [3.4, 11.0, 1.8], hdr='7030A0')
d.add_heading('鎮静か全身麻酔か ── TAVI・TEERの臨床研究', level=3)
bullets([
    'SOLVE-TAVI（RCT、447例）：意識下鎮静と全身麻酔で、30日の複合エンドポイント（全死亡、脳卒中、心筋梗塞、抗菌薬を要する感染、急性腎障害）は27.2%対26.4%で同等だった[15]。',
    '米国TVTレジストリー（経大腿TAVR 10,997例）では、意識下鎮静で院内死亡1.6%対2.5%、30日死亡2.9%対4.1%と低く、鎮静から全身麻酔への移行は5.9%だった[16]。120,080例・559施設の操作変数解析でも、鎮静で30日死亡が0.5%低く、在院日数が0.8日短かった[17]。いずれも観察研究で、鎮静を選べる患者が選ばれている点に注意がいる。',
    'MitraClip 232例では、深鎮静と全身麻酔で安全性の複合エンドポイントに差がなく、深鎮静から全身麻酔への移行は2%、準備時間とICU滞在は深鎮静で短かった[18]。国内の適正使用指針は、いずれの治療でも全身麻酔ができる体制を施設要件にしている[10,12,14]。',
])
d.add_heading('TEEの禁忌（表12、p.19）[3]', level=3)
para('絶対禁忌：食道腫瘍、傍食道型食道裂孔ヘルニア、活動性消化管出血、食道狭窄・憩室、食道穿孔・裂傷、内臓穿孔。相対禁忌：胸部・頚部の放射線治療歴、消化管術後、最近の上部消化管出血、バレット食道、嚥下障害の既往、頚椎可動性制限、症状のある食道裂孔ヘルニア、食道静脈瘤、凝固障害・血小板減少、活動性食道炎・消化性潰瘍。術中TEEは通常より偶発症のリスクが高く、強い圧迫やロック下の操作を避け、先端温度に注意する。SHDのTEE担当医の要件として、日本超音波医学会専門医、JB-POT、SHD心エコー図認証医などが示されている（表13）。')
para('弁膜症の評価では、TEE時の鎮静薬や全身麻酔・術中の血流低下が逆流を過小評価させるため、術中エコーは逆流の重症度評価には不適切とされる（p.27）[4]。')

# ================= references
d.add_heading('参考文献', level=1)
refs = [
    '日本循環器学会. 2022年改訂版 非心臓手術における合併心疾患の評価と管理に関するガイドライン. 2022. https://www.j-circ.or.jp/cms/wp-content/uploads/2022/03/JCS2022_hiraoka.pdf',
    '日本循環器学会／日本心臓血管外科学会／日本心臓病学会／日本心血管インターベンション治療学会. 2023年 フォーカスアップデート版 PCPS/ECMO/循環補助用心内留置型ポンプカテーテルの適応・操作. 2023. https://www.j-circ.or.jp/cms/wp-content/uploads/2023/03/JCS2023_nishimura.pdf',
    '日本循環器学会ほか. 2021年改訂版 先天性心疾患、心臓大血管の構造的疾患（structural heart disease）に対するカテーテル治療のガイドライン. 2021. https://www.j-circ.or.jp/cms/wp-content/uploads/2021/03/JCS2021_Sakamoto_Kawamura.pdf',
    '日本循環器学会／日本胸部外科学会／日本血管外科学会／日本心臓血管外科学会. 2020年改訂版 弁膜症治療のガイドライン. 2020. https://www.j-circ.or.jp/cms/wp-content/uploads/2020/04/JCS2020_Izumi_Eishi.pdf',
    '日本循環器学会／日本不整脈心電学会. 2021年 フォーカスアップデート版 不整脈非薬物治療. 2021. https://www.j-circ.or.jp/cms/wp-content/uploads/2021/03/JCS2021_Kurita_Nogami.pdf',
    '日本循環器学会／日本産科婦人科学会. 2026年 フォーカスアップデート版 心血管疾患患者の妊娠・出産の適応と診療. 2026. https://www.j-circ.or.jp/cms/wp-content/uploads/2026/03/JCS2026_Kamiya_Katsuragi.pdf',
    '日本循環器学会／日本小児循環器学会／日本成人先天性心疾患学会. 2025年改訂版 成人先天性心疾患診療ガイドライン. 2025. https://www.j-circ.or.jp/cms/wp-content/uploads/2025/03/JCS2025_Yamagishi.pdf',
    '日本循環器学会／日本心臓血管外科学会／日本胸部外科学会／日本血管外科学会. 2021年改訂版 重症心不全に対する植込型補助人工心臓治療ガイドライン. 2021. https://www.j-circ.or.jp/cms/wp-content/uploads/2021/03/JCS2021_Ono_Yamaguchi.pdf',
    '補助人工心臓治療関連学会協議会 インペラ部会. IMPELLA 適正使用指針（第6版）. 2024年6月18日改定. https://j-pvad.jp/guidance/',
    '日本心血管インターベンション治療学会ほか. 経皮的心房中隔欠損閉鎖セットに関する適正使用指針. 2024年11月1日作成、2025年12月1日改訂. http://www.cvit.jp/docs/about/committee/device/asd_pda/guideline-occluder-atrial_septum.pdf',
    '経カテーテル的心臓弁治療関連学会協議会. 経カテーテル的大動脈弁置換術 実施施設基準. 2013年7月18日施行、2026年6月8日改定. http://j-tavr.com/guideline.html',
    '日本循環器学会ほか. 弁尖間クリッピング式の経皮的僧帽弁接合不全修復システムに関する適正使用指針. 2017年11月18日初版、2024年12月6日改訂. https://www.j-circ.or.jp/cms/wp-content/uploads/2024/12/MTEERusage-guidelines20241206.pdf',
    '日本循環器学会ほか. 弁尖間クリッピング式の経皮的三尖弁接合不全修復システムに関する適正使用指針. 2025年9月10日初版. https://www.j-circ.or.jp/cms/wp-content/uploads/2025/12/TriClipusage-guidelines20250910_20251208.pdf',
    '日本循環器学会ほか. 左心耳閉鎖システムに関する適正使用指針. 2024年12月6日改訂. https://www.j-circ.or.jp/cms/wp-content/uploads/2024/12/WMusage-guidelines20241206.pdf',
    'Thiele H, Kurz T, Feistritzer HJ, et al. General versus local anesthesia with conscious sedation in transcatheter aortic valve implantation: the randomized SOLVE-TAVI trial. Circulation. 2020;142(15):1437-1447. doi:10.1161/CIRCULATIONAHA.120.046451',
    'Hyman MC, Vemulapalli S, Szeto WY, et al. Conscious sedation versus general anesthesia for transcatheter aortic valve replacement: insights from the National Cardiovascular Data Registry Society of Thoracic Surgeons/American College of Cardiology Transcatheter Valve Therapy Registry. Circulation. 2017;136(22):2132-2140. doi:10.1161/CIRCULATIONAHA.116.026656',
    'Butala NM, Chung M, Secemsky EA, et al. Conscious sedation versus general anesthesia for transcatheter aortic valve replacement: variation in practice and outcomes. JACC Cardiovasc Interv. 2020;13(11):1277-1287. doi:10.1016/j.jcin.2020.03.008',
    'Horn P, Hellhammer K, Minier M, et al. Deep sedation vs. general anesthesia in 232 patients undergoing percutaneous mitral valve repair using the MitraClip system. Catheter Cardiovasc Interv. 2017;90(7):1212-1219. doi:10.1002/ccd.26884',
]
for i, r in enumerate(refs, 1):
    p = d.add_paragraph(f'{i}. {r}'); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.first_line_indent = Cm(-0.5)
    for rr in p.runs: rr.font.size = Pt(8.5)

# page numbers in footer
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for tag, txt in [('begin', None), (None, 'PAGE'), ('end', None)]:
    r = fp.add_run()
    if tag:
        e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), tag); r._r.append(e)
    else:
        e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; r._r.append(e)

d.save('guideline_periop.docx')
print('ok')
